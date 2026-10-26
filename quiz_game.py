import html
import io
import json
import random
import math
import re
import time
import uuid
from pathlib import Path
from nicegui import app, ui

try:
    import segno          # QR코드 만드는 도구 (pip install segno)
except ImportError:
    segno = None

# =====================================================================
#  기본 설정 (여기 숫자나 글자를 바꾸면 게임 규칙이 바뀌어요)
# =====================================================================
JOIN_URL = ""                  # 입장 QR에 넣을 주소. 비워두면 진행자 화면 주소를 보고 자동으로 정해요
SIGN_SIZE = "크게"             # 대기실 QR 팻말 크기: "크게"(판자 1.5배, 멀리서도 잘 찍혀요) 또는 "보통"
HOST_PASSWORD = "quiz1234"     # 진행자 비밀번호 (꼭 나만 아는 걸로 바꾸세요!)
TIME_LIMIT = 15                # 문제당 제한 시간(초)
STREAK_FROM = 3                # 몇 문제 연속으로 맞히면 보너스를 줄지
STREAK_BONUS = 200             # 연속 정답 보너스 점수
COLORS = ["#E21B3C", "#1368CE", "#D89E00", "#26890C"]   # 4지선다 1~4번 색깔
OX_COLORS = ["#1368CE", "#E21B3C"]                       # O, X 색깔
CARD_IMAGES = ["card1_red.jpg", "card2_blue.jpg", "card3_yellow.jpg", "card4_green.jpg"]  # 1~4번 카드 그림
OX_CARD_IMAGES = ["card2_blue.jpg", "card1_red.jpg"]     # O, X 카드 그림 (images/cards 폴더 안)
TEAM_SCORE = "평균"           # 팀 점수 계산: "평균" 또는 "합계"
TEAM_COLORS = ["#E21B3C", "#1368CE", "#D89E00", "#26890C", "#8E44AD", "#16A085", "#E67E22", "#2C3E50"]
FINAL_PHRASE = ["오늘의", "퀴즈 챔피언", "탄생!"]    # 시상식 뒤에 크게 깔리는 문구 (줄마다 하나씩)
FINAL_PHRASE_TEAM = ["최고의", "팀워크", "우승 팀!"]   # 팀 시상식 뒤에 깔리는 문구
FINAL_PHRASE_COLOR = "#8E2A1F"                     # 그 문구 색깔
CONFETTI_COLORS = ["#E21B3C", "#1368CE", "#D89E00", "#26890C", "#F5A33A", "#FFFFFF"]   # 꽃가루 색깔
LOBBY_MUSIC = "music1.mp3"     # 대기실, 순위, 최종 결과 때 나오는 음악 (sounds 폴더 안)
QUIZ_MUSIC = "music2.mp3"      # 문제 푸는 동안 나오는 음악 (sounds 폴더 안)
MUSIC_VOLUME = 0.45            # 배경음악 크기 (0 ~ 1)
REVEAL_SOUND = "reveal.mp3"    # 정답 공개 효과음 (sounds 폴더 안)
REVEAL_VOLUME = 0.9            # 정답 공개 효과음 크기 (0 ~ 1)
BOARD_TOP_HOST = 8             # 진행자 큰 화면 점수판에 보여줄 상위 인원
BOARD_TOP_PLAYER = 3           # 휴대폰 점수판에 보여줄 상위 인원 (+ 내 순위)

# =====================================================================
#  폴더 준비
# =====================================================================
BASE = Path(__file__).parent
QUIZ_DIR = BASE / "quizzes"             # 퀴즈 세트 공책들을 보관하는 폴더
IMAGE_DIR = BASE / "images"             # 캐릭터 그림 폴더
PHOTO_DIR = IMAGE_DIR / "questions"     # 문제에 넣은 사진 폴더
CARD_DIR = IMAGE_DIR / "cards"          # 보기 카드 그림 폴더
SOUND_DIR = BASE / "sounds"             # 배경음악 폴더
QUIZ_DIR.mkdir(exist_ok=True)
PHOTO_DIR.mkdir(parents=True, exist_ok=True)
SOUND_DIR.mkdir(exist_ok=True)
app.add_static_files("/images", IMAGE_DIR)
app.add_static_files("/sounds", SOUND_DIR)

DEFAULT_QUESTIONS = [
    {"유형": "4지선다", "문제": "대한민국의 수도는?", "보기": ["부산", "서울", "대구", "광주"], "정답": 2, "사진": ""},
    {"유형": "OX", "문제": "무지개는 7가지 색이다.", "보기": ["O", "X"], "정답": 1, "사진": ""},
    {"유형": "단답형", "문제": "1 + 1 은? (숫자로 쓰세요)", "보기": [], "정답": ["2"], "사진": ""},
]

# =====================================================================
#  퀴즈 세트 (공책) 다루기
# =====================================================================
def clean_name(name):
    """세트 이름에서 파일 이름에 못 쓰는 글자 빼기"""
    return re.sub(r"[^\w\s-]", "", name or "").strip()


def set_path(name):
    return QUIZ_DIR / f"{name}.json"


def list_sets():
    return sorted(p.stem for p in QUIZ_DIR.glob("*.json"))


def fix_question(q):
    """예전 형식 문제도 쓸 수 있게 빠진 칸 채우기"""
    q.setdefault("유형", "4지선다")
    q.setdefault("보기", [])
    q.setdefault("사진", "")
    return q


def load_set(name):
    try:
        data = json.loads(set_path(name).read_text(encoding="utf-8"))
    except Exception:
        data = []
    return [fix_question(q) for q in data]


def save_set():
    set_path(game["세트"]).write_text(
        json.dumps(questions, ensure_ascii=False, indent=2), encoding="utf-8")


# 처음 실행할 때 세트가 하나도 없으면, 예전 문제 공책(questions.json)을 '기본 퀴즈'로 옮겨요
if not list_sets():
    try:
        first = json.loads((BASE / "questions.json").read_text(encoding="utf-8"))
    except Exception:
        first = DEFAULT_QUESTIONS
    set_path("기본 퀴즈").write_text(
        json.dumps([fix_question(dict(q)) for q in first], ensure_ascii=False, indent=2),
        encoding="utf-8")

# =====================================================================
#  메모장 글로 문제 한꺼번에 만들기
#  (문제 만들기 화면에 붙여넣거나 .txt 파일을 올리면 이 코드가 읽어서 문제로 바꿔요)
# =====================================================================
SAMPLE_TEXT = """# 이렇게 # 로 시작하는 줄은 메모용이라 무시돼요
# 문제와 문제 사이는 빈 줄로 띄워주세요

Q: 대한민국의 수도는?
1) 부산
2) 서울
3) 대구
4) 광주
정답: 2

Q: 무지개는 7가지 색이다.
정답: O

Q: 1 + 1 은? (숫자로 쓰세요)
정답: 2, 둘

Q: 크레인 아래에서 작업할 때 꼭 써야 하는 보호구는?
A. 선글라스
B. 안전모
C. 귀마개
D. 장갑
정답: B
"""

_ANSWER_RE = re.compile(r"^\s*(?:정답|answer|ans)\s*[:：=]?\s*(.+?)\s*$|^\s*답\s*[:：=]\s*(.+?)\s*$", re.I)
_QPREFIX_RE = re.compile(r"^\s*(?:Q|문제)\s*\d*\s*[:.)：]\s*", re.I)
_QNUM_RE = re.compile(r"^\s*\d+\s*[.)]\s+")
_CHOICE_RE = re.compile(r"^\s*(?:([1-4])|([A-Da-d]))\s*[).:．、]\s*(.+?)\s*$")
_CIRCLED_RE = re.compile(r"^\s*([①②③④])\s*[).:]?\s*(.+?)\s*$")
_CIRCLED = {"①": 1, "②": 2, "③": 3, "④": 4}
_OX_TRUE = {"o", "ⓞ", "⭕", "○", "ㅇ", "맞다", "맞음", "맞아요", "참", "true", "t", "yes", "y"}
_OX_FALSE = {"x", "ⓧ", "❌", "×", "✕", "✗", "틀리다", "틀림", "틀려요", "거짓", "false", "f", "no", "n"}


def decode_text(data):
    """메모장 파일 글자 읽기 (윈도우 메모장은 UTF-8 또는 ANSI(cp949)로 저장해요)"""
    for enc in ("utf-8-sig", "cp949"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def _norm_ox(text):
    return re.sub(r"[\s.!]", "", str(text)).lower()


def _match_choice(line):
    m = _CIRCLED_RE.match(line)
    if m:
        return m.group(2)
    m = _CHOICE_RE.match(line)
    return m.group(3) if m else None


def _split_blocks(text):
    """글 전체를 문제 한 개씩 덩어리로 나누기"""
    blocks, cur = [], []

    def flush():
        nonlocal cur
        if cur:
            blocks.append(cur)
        cur = []

    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw.strip().lstrip("\ufeff")
        if not line or re.fullmatch(r"[-=_*~]{3,}", line):
            flush()
            continue
        if line.startswith("#") or line.startswith("//"):
            continue
        if re.fullmatch(r"\[\s*(4지선다|OX|ox|단답형?|주관식|객관식)\s*\]", line):
            continue
        if _QPREFIX_RE.match(line) and cur:     # "Q:" 로 시작하면 새 문제
            flush()
        cur.append(line)
    flush()
    return blocks


def _parse_block(lines):
    """덩어리 하나 → 문제 하나. (문제, None) 또는 (None, 오류 설명)"""
    q_lines, choices, answer_raw = [], [], None
    for line in lines:
        m = _ANSWER_RE.match(line)
        if m:
            answer_raw = (m.group(1) or m.group(2) or "").strip()
            continue
        if q_lines:
            c = _match_choice(line)
            if c:
                choices.append(c)
                continue
            if choices:                      # 보기가 길어서 다음 줄로 넘어간 경우
                choices[-1] += " " + line
                continue
        q_lines.append(line)

    if not q_lines:
        return None, "문제 글이 없어요"
    text = " ".join(q_lines)
    text = _QPREFIX_RE.sub("", text, count=1) if _QPREFIX_RE.match(text) else _QNUM_RE.sub("", text, count=1)
    text = text.strip()
    if not text:
        return None, "문제 글이 없어요"
    if not answer_raw:
        return None, "'정답:' 줄이 없어요"

    n = len(choices)
    ans_ox = _norm_ox(answer_raw)

    # --- OX ---
    ox_choices = n == 2 and all(_norm_ox(c) in _OX_TRUE | _OX_FALSE for c in choices)
    if ox_choices or (n == 0 and ans_ox in _OX_TRUE | _OX_FALSE):
        if ans_ox in _OX_TRUE:
            right = 1
        elif ans_ox in _OX_FALSE:
            right = 2
        else:
            return None, f"OX 문제의 정답은 O 또는 X로 써주세요 (지금: {answer_raw})"
        return {"유형": "OX", "문제": text, "보기": ["O", "X"], "정답": right, "사진": ""}, None

    # --- 단답형 ---
    if n == 0:
        answers = [a.strip() for a in re.split(r"[,，、|]", answer_raw) if a.strip()]
        if not answers:
            return None, "정답이 비어 있어요"
        return {"유형": "단답형", "문제": text, "보기": [], "정답": answers, "사진": ""}, None

    # --- 4지선다 ---
    if n != 4:
        return None, f"보기는 4개여야 해요 (지금 {n}개)"
    a = answer_raw.strip()
    right = None
    m = re.fullmatch(r"([1-4])\s*번?\s*[).]?", a)
    if m:
        right = int(m.group(1))
    elif a[:1] in _CIRCLED and len(a) <= 2:
        right = _CIRCLED[a[0]]
    elif re.fullmatch(r"[A-Da-d]\s*[).]?", a):
        right = "abcd".index(a[0].lower()) + 1
    else:
        for i, c in enumerate(choices, 1):          # 정답을 보기 글자 그대로 쓴 경우
            if re.sub(r"\s+", "", c).lower() == re.sub(r"\s+", "", a).lower():
                right = i
                break
    if right is None:
        return None, f"정답 '{answer_raw}' 을(를) 보기에서 찾을 수 없어요 (1~4 또는 A~D 로 써주세요)"
    return {"유형": "4지선다", "문제": text, "보기": choices, "정답": right, "사진": ""}, None


def parse_quiz_text(text):
    """메모장 글 전체 → (문제 목록, 오류 목록)"""
    good, errors = [], []
    for i, block in enumerate(_split_blocks(text), 1):
        q, err = _parse_block(block)
        if q:
            good.append(q)
        else:
            snippet = _QPREFIX_RE.sub("", block[0], count=1)[:24]
            errors.append(f"{i}번째 덩어리 ('{snippet}…'): {err}")
    return good, errors


# =====================================================================
#  교실 칠판: 모두가 같이 보는 게임 정보
# =====================================================================
game = {"단계": "대기", "번호": 0, "마감시각": 0, "세트": list_sets()[0]}
questions = load_set(game["세트"])
players = {}      # {"닉네임": {"점수", "답", "이번점수", "보너스", "연속", "기록"}}
kicked = set()    # 진행자가 내보낸 닉네임들

# ----- 팀 (진행자가 만든 팀 목록. 하나도 없으면 개인전) -----
TEAM_FILE = BASE / "teams.json"
try:
    teams = json.loads(TEAM_FILE.read_text(encoding="utf-8"))
except Exception:
    teams = []


def save_teams():
    TEAM_FILE.write_text(json.dumps(teams, ensure_ascii=False), encoding="utf-8")


def team_color(team):
    return TEAM_COLORS[teams.index(team) % len(TEAM_COLORS)] if team in teams else "#8C8C8C"


def add_team(name):
    name = clean_name(name)
    if game["단계"] != "대기":
        ui.notify("게임 중에는 팀을 바꿀 수 없어요")
    elif not name:
        ui.notify("팀 이름을 입력해주세요 (특수문자는 빼주세요)")
    elif name in teams:
        ui.notify("이미 있는 팀 이름이에요")
    else:
        teams.append(name)
        save_teams()
        return True
    return False


def remove_team(name):
    if game["단계"] != "대기":
        ui.notify("게임 중에는 팀을 바꿀 수 없어요")
        return
    if name in teams:
        teams.remove(name)
        save_teams()
        for p in players.values():          # 그 팀이었던 사람은 팀을 다시 골라야 해요
            if p["팀"] == name:
                p["팀"] = None


def team_rows(field="점수"):
    """팀별 점수 [(팀, 점수, 인원)] 높은 순서로. field가 '이번점수'면 이번 라운드 점수"""
    rows = []
    for t in teams:
        members = [p for p in players.values() if p["팀"] == t]
        if not members:
            continue
        total = sum(p[field] for p in members)
        score = round(total / len(members)) if TEAM_SCORE == "평균" else total
        rows.append((t, score, len(members)))
    return sorted(rows, key=lambda r: r[1], reverse=True)


def new_player(team=None):
    return {"점수": 0, "답": None, "이번점수": 0, "보너스": 0, "연속": 0, "기록": [], "팀": team}


def choose_set(name):
    """플레이할(그리고 편집할) 세트 바꾸기"""
    if game["단계"] != "대기":
        ui.notify("게임 중에는 세트를 바꿀 수 없어요")
        return False
    if name not in list_sets():
        return False
    game["세트"] = name
    questions[:] = load_set(name)
    return True


# =====================================================================
#  도우미
# =====================================================================
def time_left():
    return max(0, math.ceil(game["마감시각"] - time.time()))


def ranking():
    return sorted(players.items(), key=lambda item: item[1]["점수"], reverse=True)


def normalize(text):
    """띄어쓰기 없애고 영어는 소문자로 (단답형 채점용)"""
    return re.sub(r"\s+", "", str(text)).lower()


def is_correct(q, ans):
    if ans is None:
        return False
    if q["유형"] == "단답형":
        return normalize(ans) in [normalize(a) for a in q["정답"]]
    return ans == q["정답"]


def answer_text(q):
    if q["유형"] == "단답형":
        return " 또는 ".join(q["정답"])
    if q["유형"] == "OX":
        return q["보기"][q["정답"] - 1]
    return f"{q['정답']}번, {q['보기'][q['정답'] - 1]}"


def show_answer(q, ans):
    """참가자가 고른 답을 글자로 (엑셀용)"""
    if ans is None:
        return "(무응답)"
    if q["유형"] == "단답형":
        return str(ans)
    return q["보기"][ans - 1]


def photo_html(q, css_class):
    if not q.get("사진"):
        return ""
    return f'<img src="/images/questions/{html.escape(q["사진"])}" class="{css_class}">'


# =====================================================================
#  게임 진행
# =====================================================================
def start_game():
    if not questions:
        ui.notify("이 세트에는 문제가 없어요. 먼저 문제를 만들어주세요")
        return
    if teams:
        missing = [n for n, p in players.items() if p["팀"] not in teams]
        if missing:
            ui.notify(f"아직 팀을 안 고른 사람이 있어요: {', '.join(missing)}", type="warning")
            return
    game["번호"] = 0
    for p in players.values():
        p.update(점수=0, 연속=0, 기록=[])
    start_question()


def start_question():
    game["단계"] = "문제"
    game["마감시각"] = time.time() + TIME_LIMIT
    for p in players.values():
        p.update(답=None, 이번점수=0, 보너스=0)


def go_to_answer():
    """정답 공개로 넘어가면서 각자 기록 남기기"""
    if game["단계"] != "문제":
        return
    q = questions[game["번호"]]
    for p in players.values():
        right = is_correct(q, p["답"])
        if not right:
            p["연속"] = 0
        p["기록"].append({"답": p["답"], "맞음": right, "점수": p["이번점수"]})
    game["단계"] = "정답"


def check_time():
    """시간이 다 됐거나 모두 답했으면 자동으로 정답 공개"""
    if game["단계"] == "문제":
        everyone = len(players) > 0 and all(p["답"] is not None for p in players.values())
        if time_left() == 0 or everyone:
            go_to_answer()


def show_ranking():
    """이번 라운드 점수 보기"""
    game["단계"] = "순위"


def show_total():
    """누적 순위 보기"""
    game["단계"] = "누적"


def show_final_board():
    """최종 결과에서 전체 순위표 보기"""
    game["단계"] = "최종표"


def back_to_ceremony():
    """개인 시상식으로 (전체 순위표나 팀 시상식에서 돌아오기)"""
    game["단계"] = "끝"


def show_team_ceremony():
    """팀 시상식 보기"""
    game["단계"] = "팀끝"


def next_question():
    game["번호"] += 1
    if game["번호"] >= len(questions):
        game["단계"] = "끝"
    else:
        start_question()


def reset_game():
    game.update({"단계": "대기", "번호": 0, "마감시각": 0})
    players.clear()
    kicked.clear()


def kick(name):
    players.pop(name, None)
    kicked.add(name)
    ui.notify(f"{name}님을 내보냈어요")


def submit_answer(name, ans):
    """참가자가 답을 냈을 때 채점 (빨리 맞힐수록 + 연속 정답 보너스)"""
    p = players.get(name)
    if p is None or game["단계"] != "문제" or p["답"] is not None:
        return
    if isinstance(ans, str):
        ans = ans.strip()
        if not ans:
            ui.notify("답을 입력해주세요")
            return
    p["답"] = ans
    if is_correct(questions[game["번호"]], ans):
        remain = max(0, game["마감시각"] - time.time())
        p["연속"] += 1
        p["보너스"] = STREAK_BONUS if p["연속"] >= STREAK_FROM else 0
        p["이번점수"] = 500 + int(500 * remain / TIME_LIMIT) + p["보너스"]
        p["점수"] += p["이번점수"]


# =====================================================================
#  보기별 선택 인원 막대그래프 (단답형 정답 공개용)
# =====================================================================
def bar_chart_html(q):
    items = []   # (글자, 인원, 색깔, 정답인지)
    if q["유형"] in ("4지선다", "OX"):
        colors = COLORS if q["유형"] == "4지선다" else OX_COLORS
        for n, choice in enumerate(q["보기"], 1):
            count = sum(1 for p in players.values() if p["답"] == n)
            items.append((choice, count, colors[n - 1], n == q["정답"]))
    else:  # 단답형: 많이 나온 답 5개
        counts, first_seen = {}, {}
        for p in players.values():
            if p["답"] is not None:
                key = normalize(p["답"])
                counts[key] = counts.get(key, 0) + 1
                first_seen.setdefault(key, p["답"])
        top = sorted(counts.items(), key=lambda kv: kv[1], reverse=True)[:5]
        for key, count in top:
            right = is_correct(q, first_seen[key])
            items.append((first_seen[key], count, "#26890C" if right else "#8C8C8C", right))
    no_answer = sum(1 for p in players.values() if p["답"] is None)
    if no_answer:
        items.append(("무응답", no_answer, "#B5B5B5", False))
    if not items:
        return '<div class="bars-empty">아무도 답하지 않았어요</div>'

    biggest = max(count for _, count, _, _ in items) or 1
    cols = []
    for label, count, color, right in items:
        height = 4 + int(62 * count / biggest)          # 화면 높이에 맞춰 % 로
        mark = "✅ " if right else ""
        dim = "" if right else " dim"
        cols.append(
            f'<div class="bar-col{dim}"><div class="bar-num">{count}명</div>'
            f'<div class="bar" style="height:{height}%; background:{color}"></div>'
            f'<div class="bar-label">{mark}{html.escape(str(label))}</div></div>')
    return '<div class="bars">' + "".join(cols) + "</div>"


# =====================================================================
#  결과 엑셀 만들기
# =====================================================================
def make_excel():
    from openpyxl import Workbook   # pip install openpyxl 이 필요해요
    wb = Workbook()

    ws = wb.active
    ws.title = "순위"
    ws.append(["순위", "닉네임", "팀", "총점", "맞힌 문제 수"])
    for rank, (name, p) in enumerate(ranking(), 1):
        ws.append([rank, name, p["팀"] or "", p["점수"], sum(1 for r in p["기록"] if r["맞음"])])

    if teams:
        ws_team = wb.create_sheet("팀 순위")
        ws_team.append(["순위", "팀", f"팀 점수({TEAM_SCORE})", "인원"])
        for rank, (t, score, count) in enumerate(team_rows(), 1):
            ws_team.append([rank, t, score, count])

    ws2 = wb.create_sheet("문제별 정답률")
    ws2.append(["번호", "유형", "문제", "정답", "정답자 수", "참가자 수", "정답률(%)"])
    for i, q in enumerate(questions):
        records = [p["기록"][i] for p in players.values() if len(p["기록"]) > i]
        right = sum(1 for r in records if r["맞음"])
        rate = round(right / len(records) * 100) if records else 0
        ws2.append([i + 1, q["유형"], q["문제"], answer_text(q), right, len(records), rate])

    ws3 = wb.create_sheet("참가자별 답")
    ws3.append(["닉네임"] + [f"{i + 1}번" for i in range(len(questions))])
    for name, p in ranking():
        row = [name]
        for i, q in enumerate(questions):
            if i < len(p["기록"]):
                r = p["기록"][i]
                row.append(show_answer(q, r["답"]) + (" ✅" if r["맞음"] else " ❌"))
            else:
                row.append("")
        ws3.append(row)

    for sheet in wb.worksheets:
        for col in sheet.columns:
            width = max(len(str(c.value or "")) for c in col) * 2 + 4
            sheet.column_dimensions[col[0].column_letter].width = min(width, 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def download_excel():
    try:
        data = make_excel()
    except ImportError:
        ui.notify("엑셀 기능을 쓰려면 명령 프롬프트에서 pip install openpyxl 을 먼저 해주세요",
                  type="warning", timeout=8000)
        return
    ui.download.content(data, f"quiz_result_{time.strftime('%Y%m%d_%H%M')}.xlsx")


# =====================================================================
#  디자인 (화면 전체에 입히는 단체복)
#  - 모든 게임 화면은 "지금 보는 화면 크기" 한 장 안에 맞춰져요 (스크롤 없음)
#  - 글자 크기는 화면의 가로/세로 중 좁은 쪽(vmin)을 기준으로 정해요
# =====================================================================
ui.add_head_html("""
<meta name="color-scheme" content="light">
<link href="https://fonts.googleapis.com/css2?family=Black+Han+Sans&family=Jua&display=swap" rel="stylesheet">
<style>
  :root { --pink: #E8A598; --cream: #F4EFEA; --ink: #2B2523; color-scheme: light; }
  html { scrollbar-gutter: stable; }
  body { background-color: var(--pink); font-family: 'Jua', sans-serif !important; color: var(--ink);
         touch-action: manipulation; -webkit-text-size-adjust: 100%; }
  .nicegui-content { padding: 0 !important; gap: 0 !important; }
  input, textarea { font-size: 16px !important; }      /* 아이폰에서 입력칸 누를 때 자동 확대 막기 */

  /* ===== 한 화면 틀: 지금 보는 화면 크기에 딱 맞게, 가운데 정렬 ===== */
  .screen { width: 100%; height: 100vh; height: 100dvh; box-sizing: border-box; overflow: hidden;
            padding: calc(env(safe-area-inset-top, 0px) + 1.6vmin) 1.6vmin
                     calc(env(safe-area-inset-bottom, 0px) + 1.6vmin);
            display: flex; flex-direction: column; align-items: center; justify-content: center; gap: 1.4vmin; }
  .screen > * { flex: 0 0 auto; max-width: 1200px; }
  .actions { display: flex; flex-wrap: wrap; justify-content: center; gap: 1vmin; }
  .actions .q-btn { font-family: 'Jua', sans-serif; font-size: clamp(14px, 2.2vmin, 20px); }
  .msg-big { font-size: clamp(22px, 5.5vmin, 52px); text-align: center; word-break: keep-all; }
  .msg { font-size: clamp(15px, 2.8vmin, 24px); text-align: center; word-break: keep-all; }

  /* ===== 대기실 ===== */
  .topbar { width: 100%; max-width: 1500px !important; background: var(--cream); border-radius: 1.6vmin; padding: 1vmin 2vmin;
            display: flex; flex-wrap: wrap; align-items: center; justify-content: space-between; gap: 1vmin; }
  .menu { font-size: clamp(14px, 2.2vmin, 20px); letter-spacing: 0.04em; }
  .brand { font-family: 'Black Han Sans', sans-serif; font-size: clamp(22px, 4vmin, 36px); }
  .menu-btn { font-family: 'Jua', sans-serif; font-size: clamp(14px, 2vmin, 18px); border-radius: 8px; }
  .hero { flex: 1 1 auto !important; min-height: 0; width: 100%; position: relative; overflow: hidden;
          display: flex; flex-direction: column; align-items: center; justify-content: center; }
  .hero-title { font-family: 'Black Han Sans', sans-serif; color: var(--cream); line-height: 1;
                font-size: min(12vw, 15dvh, 200px); position: relative; z-index: 1; }
  .hero-char { height: min(32dvh, 420px); width: auto; margin-top: max(-6vw, -7dvh); position: relative;
               z-index: 2; animation: bob 3s ease-in-out infinite; filter: drop-shadow(0 12px 12px rgba(0,0,0,0.15)); }
  @keyframes bob { 0%, 100% { transform: translateY(0); } 50% { transform: translateY(-1.4vmin); } }
  .bit { position: absolute; border-radius: 50%; z-index: 0; animation: float 6s ease-in-out infinite; }
  @keyframes float { 0%, 100% { transform: translateY(0) rotate(-20deg); }
                     50% { transform: translateY(-22px) rotate(15deg); } }
  .lobby-names { width: 100%; max-height: 26dvh; overflow-y: auto; }

  /* ===== 입장 QR: 용강이가 낑낑대며 들고 있는 안내 팻말 =====
     크기는 모두 --w(용강이+팻말 전체 폭)를 기준으로 계산해요 */
  .hero.has-qr { flex-direction: row; gap: 5vw; }
  .hero-titles { display: flex; flex-direction: column; align-items: center; position: relative; z-index: 1; }
  .hero.has-qr .hero-title { font-size: min(10vw, 15dvh, 190px); }
  .qr-mascot { --w: min(36dvh, 30vw); --body: #F5A22C; --line: #3B3C8E; position: relative; z-index: 2; flex: none;
               width: var(--w); height: calc(var(--w) * 1.58); cursor: zoom-in; }
  .qm-char { position: absolute; left: 50%; top: 0; width: 62%; transform: translateX(-50%);
             animation: qm-head 0.9s ease-in-out infinite; }
  .qm-foot { position: absolute; top: calc(var(--w) * 1.44); width: 17%; height: calc(var(--w) * 0.13);
             background: var(--body); box-sizing: border-box; border: calc(var(--w) * 0.014) solid var(--line);
             border-radius: 50% 50% 45% 45%; animation: qm-shake 0.18s linear infinite; }
  .qm-foot.l { left: 30%; }
  .qm-foot.r { right: 30%; animation-delay: 0.09s; }
  .qm-sign { position: absolute; left: 3%; right: 3%; top: calc(var(--w) * 0.55);
             animation: qm-sign 0.9s ease-in-out infinite; transform-origin: 50% 100%; }
  .qm-board { background: #FFFFFF; border: calc(var(--w) * 0.035) solid #B87333; border-radius: calc(var(--w) * 0.05);
              padding: calc(var(--w) * 0.025); box-sizing: border-box;
              box-shadow: 0 calc(var(--w) * 0.03) calc(var(--w) * 0.05) rgba(0,0,0,0.25),
                          inset 0 0 0 calc(var(--w) * 0.008) #8A5523; }
  .qm-board svg { width: 100%; height: auto; display: block; }
  .qm-board .qm-msg { aspect-ratio: 1; display: flex; align-items: center; justify-content: center; text-align: center;
                      font-size: calc(var(--w) * 0.07); line-height: 1.4; }
  .qm-hand { position: absolute; top: calc(var(--w) * 0.49); width: 15%; aspect-ratio: 1 / 0.9; background: var(--body);
             border: calc(var(--w) * 0.014) solid var(--line); border-radius: 50%; box-sizing: border-box;
             animation: qm-sign 0.9s ease-in-out infinite; }
  .qm-hand.l { left: -3%; }
  .qm-hand.r { right: -3%; }
  .qm-sweat { position: absolute; font-size: calc(var(--w) * 0.11); animation: qm-drip 1.4s ease-in infinite; }
  .qm-sweat.s1 { left: 12%; top: 6%; }
  .qm-sweat.s2 { right: 14%; top: 14%; animation-delay: 0.7s; }
  @keyframes qm-sign { 0%, 100% { transform: rotate(-2.5deg) translateY(0); } 50% { transform: rotate(2.5deg) translateY(2%); } }
  @keyframes qm-head { 0%, 100% { transform: translateX(-50%) translateY(0); } 50% { transform: translateX(-50%) translateY(2.5%); } }
  @keyframes qm-shake { 0%, 100% { transform: translateX(0); } 50% { transform: translateX(6%); } }
  @keyframes qm-drip { 0% { opacity: 0; transform: translateY(-20%); } 25% { opacity: 1; }
                       100% { opacity: 0; transform: translateY(70%); } }
  /* 첨부 그림 기반: 팻말을 번쩍 들고 휘청휘청 */
  .qr-sign { position: relative; flex: none; height: min(62dvh, 60vw); max-height: 100%; cursor: zoom-in;
             transform-origin: 45% 96%; animation: qs-sway 1s ease-in-out infinite; z-index: 2; }
  .qr-sign img { width: 100%; height: 100%; display: block; }
  .qs-qr { position: absolute; aspect-ratio: 1; box-sizing: border-box; padding: 2.5%;
           transform: translate(-50%, -50%) rotate(var(--a)); background: #FFFFFF; border-radius: 4%;
           box-shadow: 0 0 0 0.25vmin rgba(33,40,121,0.35); }
  .qs-qr svg { width: 100%; height: 100%; display: block; }
  .qs-qr .qm-msg { font-size: 1.6vmin; height: 100%; }
  .qs-sweat { position: absolute; left: 14%; top: 31%; font-size: min(4.5dvh, 4vw);
              animation: qm-drip 1.3s ease-in infinite; }
  @keyframes qs-sway { 0%, 100% { transform: rotate(-1.6deg); } 50% { transform: rotate(1.6deg) scaleY(0.99); } }
  /* 원래 용강이(hero-char)와 똑같은 크기·자리의 칸.
     팻말 그림 속 용강이 몸이 원래 용강이와 같은 크기·위치에 오도록 잰 값으로 맞췄어요 */
  .hero-sign-slot { position: relative; z-index: 2; flex: none; height: min(32dvh, 420px); aspect-ratio: 563 / 832;
                    margin-top: max(-6vw, -7dvh); }
  .hero-sign-slot .qr-sign { position: absolute; height: var(--sh); max-height: none; left: var(--sl); top: var(--st); }
  @media (orientation: portrait) {
    .hero-sign-slot { height: min(24dvh, 50vw); }
    .qr-sign { height: min(42dvh, 110vw); }
    .hero.has-qr { flex-direction: column; gap: 1vmin; }
    .hero.has-qr .hero-title { font-size: min(13vw, 8dvh); }
    .qr-mascot { --w: min(30dvh, 62vw); }
  }
  .qr-big { position: fixed; inset: 0; z-index: 70; background: rgba(43,37,35,0.85); display: none;
            align-items: center; justify-content: center; cursor: zoom-out; }
  .qr-big.show { display: flex; }
  .qr-big-box { --w: min(80dvh, 88vw); width: var(--w); }
  .qr-big-title { font-family: 'Black Han Sans', sans-serif; font-size: clamp(20px, 4.5vmin, 54px); text-align: center; }
  .qr-big-url { font-family: 'Black Han Sans', sans-serif; font-size: clamp(16px, 3vmin, 34px); word-break: break-all;
                text-align: center; }
  .names { display: flex; flex-wrap: wrap; gap: 0.8vmin; justify-content: center; }
  .name-chip { display: flex; align-items: center; gap: 2px; background: var(--cream); border-radius: 999px;
               padding: 2px 6px 2px 16px; font-size: clamp(14px, 2.4vmin, 22px); }
  .team-groups { width: 100%; display: grid; gap: 1vmin; grid-template-columns: repeat(auto-fit, minmax(200px, 1fr)); }
  .team-group { background: var(--cream); border-radius: 1.4vmin; padding: 1vmin 1.2vmin; border-top: 8px solid; }
  .team-head { font-family: 'Black Han Sans', sans-serif; font-size: clamp(16px, 2.8vmin, 24px); margin-bottom: 0.6vmin; }
  .team-group .name-chip { background: white; font-size: clamp(13px, 2vmin, 18px); }

  /* ===== 입장 화면 (휴대폰) ===== */
  .join-box { width: min(100%, 480px); display: flex; flex-direction: column; align-items: stretch; gap: 1.8vmin; }
  .team-pick { width: 100%; display: grid; grid-template-columns: repeat(2, 1fr); gap: 1vmin; }
  .team-pick .q-btn { font-family: 'Jua', sans-serif; font-size: clamp(15px, 2.8vmin, 20px); padding: 1.6vmin 6px;
                      border-radius: 12px; }
  .team-dot { width: 18px; height: 18px; border-radius: 50%; flex: none; }
  .team-badge { color: white; border-radius: 999px; padding: 0.5em 1.4em; font-size: clamp(16px, 3vmin, 22px); }

  /* ===== 문제 / 정답 / 점수판이 들어가는 크림색 판 ===== */
  .quiz-stage { width: 100%; max-width: 1200px; flex: 1 1 auto !important; min-height: 0; box-sizing: border-box;
                background: var(--cream); border-radius: 2.4vmin; overflow: hidden;
                padding: 3vmin 1.6vmin 1.6vmin; display: flex; flex-direction: column; align-items: center; gap: 1.2vmin; }
  .stage-title { flex: none; width: 100%; max-height: 22dvh; overflow: hidden; text-align: center;
                 font-family: 'Black Han Sans', sans-serif; font-size: min(6.5vmin, 84px); line-height: 1.12;
                 word-break: keep-all; }
  .stage-sub { flex: none; font-size: clamp(13px, 2.1vmin, 20px); letter-spacing: 0.05em; text-align: center; }
  .stage-pill { background: var(--ink); color: var(--cream); border-radius: 10px; padding: 0.55em 1.2em;
                font-size: clamp(14px, 2.3vmin, 22px); min-width: 6.5em; text-align: center;
                font-variant-numeric: tabular-nums; }
  .q-photo { flex: none; height: min(22dvh, 260px); width: auto; max-width: 90%; object-fit: contain; border-radius: 1.6vmin; }
  .q-photo-sm { max-height: 26vh; max-width: 100%; border-radius: 12px; }

  /* 보기 카드: 남은 공간을 나눠 가져요. 세로 화면(휴대폰)은 2 x 2 */
  .choice-grid { width: 100%; flex: 1 1 auto; min-height: 0; display: grid;
                 grid-template-columns: repeat(4, 1fr); grid-template-rows: 1fr; gap: 1vmin; margin-top: 0.8vmin; }
  .choice-grid.two { grid-template-columns: repeat(2, 1fr); }
  @media (orientation: portrait) {
    .choice-grid { grid-template-columns: repeat(2, 1fr); grid-template-rows: 1fr 1fr; max-height: none; }
    .choice-grid.two { grid-template-rows: 1fr; }
    .q-photo { height: min(16dvh, 200px); }
  }
  .choice-card { position: relative; min-width: 0; min-height: 0; border-radius: 2vmin; color: white; overflow: hidden;
                 display: flex; flex-direction: column; padding: 1.8vmin; box-sizing: border-box;
                 box-shadow: 0 0.8vmin 0 rgba(0,0,0,0.18); transition: transform 0.12s, box-shadow 0.12s; }
  .choice-card:not(.static) { cursor: pointer; }
  .choice-card:not(.static):hover { transform: translateY(-4px); }
  .choice-card:not(.static):active { transform: translateY(3px) scale(0.98); box-shadow: 0 2px 0 rgba(0,0,0,0.18); }
  .card-top { flex: none; width: 100%; display: flex; justify-content: space-between; align-items: flex-start; }
  .choice-num { flex: none; width: max(28px, 5.5vmin); height: max(28px, 5.5vmin); border-radius: 50%;
                background: rgba(255,255,255,0.25); display: flex; align-items: center; justify-content: center;
                font-size: clamp(14px, 3vmin, 28px); }
  /* 보기 글자: 카드 한가운데. 카드 폭(cqi)과 높이(cqb)를 보고 크기를 정해요.
     짧은 답은 크게, 문장은 줄을 고르게 나눠서(가운데 정렬) 보여줘요. */
  .choice-card { container-type: size; }
  .choice-text { position: absolute; inset: 0; z-index: 2; box-sizing: border-box; overflow: hidden;
                 padding: calc(max(28px, 5.5vmin) + 3vmin) 7%;
                 display: flex; align-items: center; justify-content: center; text-align: center;
                 font-family: 'Black Han Sans', sans-serif; letter-spacing: 0.02em;
                 font-size: min(13cqi, 14cqb); line-height: 1.22; word-break: keep-all; overflow-wrap: anywhere;
                 text-wrap: balance; white-space: pre-line; }
  .choice-text.len-s  { font-size: min(46cqi, 34cqb); }    /* 1~3글자 */
  .choice-text.len-m  { font-size: min(22cqi, 20cqb); }    /* 4~8글자 */
  .choice-text.len-l  { font-size: min(17cqi, 15cqb); }    /* 9~16글자 */
  .choice-text.len-xl { font-size: min(14cqi, 12cqb); }    /* 17글자 이상 (문장) */
  .choice-text.ox { font-size: min(42cqi, 42cqb); padding: 6%; }
  /* ===== 그림 카드 =====
     그림 가운데에 은은한 그늘을 깔고, 그 위에 굵은 흰 글씨 + 그림자로 보기를 써요.
     카드 테두리와 번호 동그라미에 원래 색을 남겨서 색으로도 구분돼요. */
  .choice-card.has-img { background-color: var(--c); }
  .choice-card.has-img::before { content: ""; position: absolute; inset: 0; z-index: 0;
                                 background: var(--bg) center / cover no-repeat; transition: transform 0.45s ease; }
  .choice-card.has-img::after { content: ""; position: absolute; inset: 0; z-index: 0; pointer-events: none;
                                border-radius: inherit;
                                background: radial-gradient(ellipse 90% 55% at 50% 52%, rgba(0,0,0,0.62) 0%, rgba(0,0,0,0.38) 55%, rgba(0,0,0,0.12) 100%);
                                box-shadow: inset 0 0 0 0.5vmin var(--c), inset 0 -1vmin 0 var(--c); }
  .choice-card.has-img:not(.static):hover::before { transform: scale(1.07); }
  .choice-card.has-img .card-top { position: relative; z-index: 2; }
  .choice-card.has-img .choice-text { color: #FFFFFF;
      text-shadow: 0 0.06em 0 rgba(0,0,0,0.7), 0 0 0.25em rgba(0,0,0,0.85), 0 0 0.06em rgba(0,0,0,0.9); }
  .choice-card.has-img .choice-num { background: var(--c); border: 0.35vmin solid #FFFFFF; color: #FFFFFF;
                                     font-family: 'Black Han Sans', sans-serif; box-shadow: 0 0.3vmin 0.8vmin rgba(0,0,0,0.45); }
  .choice-card.has-img .choice-count { font-family: 'Black Han Sans', sans-serif; text-shadow: 0 0 1vmin rgba(0,0,0,0.85); }
  .shine { position: absolute; inset: 0; z-index: 1; pointer-events: none;
           background: linear-gradient(115deg, transparent 35%, rgba(255,255,255,0.28) 48%, transparent 60%);
           transform: translateX(-130%); animation: shine 3.6s ease-in-out infinite;
           animation-delay: calc(var(--i) * 0.3s + 0.8s); }
  @keyframes shine { 0%, 55% { transform: translateX(-130%); } 100% { transform: translateX(130%); } }
  /* 문제가 나올 때 카드가 하나씩 톡톡 올라와요 */
  .choice-card.enter { animation: card-in 0.45s cubic-bezier(0.2, 0.9, 0.3, 1.2) backwards;
                       animation-delay: calc(var(--i) * 0.07s); }
  @keyframes card-in { from { opacity: 0; transform: translateY(3vmin) scale(0.94); } }
  .short-box { width: 100%; max-width: 560px; margin: auto 0; display: flex; flex-direction: column; gap: 1.2vmin; }

  /* 정답 공개: 정답 카드만 밝게, 고른 인원만큼 하얀 막이 차올라요 */
  .choice-card.reveal { cursor: default; }
  /* 정답 공개 순간 화면 전환 (효과음 박자에 맞춰요)
     0초: 화면이 하얗게 번쩍 → 0.05초: 정답 제목 톡 → 0.35초: 오답 흐려짐 → 0.45초: 정답 카드 쾅 */
  .flash { position: fixed; inset: 0; background: #FFFFFF; pointer-events: none; z-index: 60;
           animation: flash 0.5s ease-out forwards; }
  @keyframes flash { 0% { opacity: 0.9; } 100% { opacity: 0; } }
  .reveal-stage { animation: stage-in 0.4s ease-out both; }
  @keyframes stage-in { from { transform: scale(0.96); } }
  .reveal-stage .stage-title { animation: hit 0.5s ease-out 0.05s both; }
  .choice-card.reveal.wrong { animation: dim 0.4s ease-out 0.35s both; }
  @keyframes dim { to { opacity: 0.4; filter: grayscale(0.85) brightness(0.8); } }
  .choice-card.reveal.right { animation: hit-ring 0.6s ease-out 0.45s both; z-index: 2; }
  @keyframes hit { 0% { transform: scale(1); } 40% { transform: scale(1.08); } 100% { transform: scale(1); } }
  @keyframes hit-ring {
    0%   { transform: scale(1);    box-shadow: 0 0.8vmin 0 rgba(0,0,0,0.18); }
    40%  { transform: scale(1.07); box-shadow: 0 0 0 0.6vmin var(--ink), 0 0.8vmin 0 rgba(0,0,0,0.18); }
    100% { transform: scale(1);    box-shadow: 0 0 0 0.6vmin var(--ink), 0 0.8vmin 0 rgba(0,0,0,0.18); } }
  .reveal-stage .fill { animation: grow 1s ease-out 0.45s both; }
  .reveal-stage .bar { animation: grow 0.9s ease-out 0.45s both; }
  @keyframes pop { 0% { transform: scale(0.92); } 60% { transform: scale(1.05); } 100% { transform: scale(1); } }
  .fill { position: absolute; left: 0; right: 0; bottom: 0; z-index: 1; background: rgba(255,255,255,0.28);
          transform-origin: bottom; animation: grow 1s ease-out; }
  @keyframes grow { from { transform: scaleY(0); } to { transform: scaleY(1); } }
  .on-top { position: relative; z-index: 1; }
  .contents { display: contents; }
  .choice-count { font-size: clamp(16px, 4vmin, 44px); }
  .my-pick { position: absolute; right: 1vmin; bottom: 1vmin; z-index: 3; background: var(--ink); color: var(--cream);
             border-radius: 8px; padding: 4px 10px; font-size: clamp(12px, 2vmin, 18px); }

  /* 단답형 막대그래프 */
  .bars { flex: 1 1 auto; min-height: 0; width: 100%; display: flex; align-items: stretch; justify-content: center;
          gap: 2vmin; padding: 1vmin; box-sizing: border-box; }
  .bar-col { flex: 0 1 120px; min-width: 0; display: flex; flex-direction: column; justify-content: flex-end;
             align-items: center; gap: 0.6vmin; }
  .bar { width: 100%; border-radius: 10px 10px 0 0; transform-origin: bottom; animation: grow 0.9s ease-out; }
  .bar-num { font-size: clamp(14px, 3vmin, 26px); }
  .bar-label { font-size: clamp(12px, 2.4vmin, 22px); text-align: center; word-break: keep-all;
               max-height: 3.6em; overflow: hidden; }
  .bar-col.dim { opacity: 0.45; }
  .bars-empty { font-size: clamp(16px, 3vmin, 24px); margin: auto; }

  /* ===== 점수판 (상위권 + 내 순위) ===== */
  .boards { flex: 1 1 auto !important; min-height: 0; width: 100%; max-width: 1200px; display: flex; gap: 1.4vmin; }
  .boards > .quiz-stage { flex: 1 1 0 !important; }
  @media (orientation: portrait) { .boards { flex-direction: column; } }
  .board { width: 100%; max-width: 900px; flex: 1 1 auto; min-height: 0; overflow: hidden;
           display: flex; flex-direction: column; gap: 0.9vmin; font-size: clamp(13px, 2.8vmin, 28px); }
  .board-row { flex: none; display: flex; align-items: center; gap: 1.6vmin; background: white; border-radius: 1.4vmin;
               padding: 0.5em 0.8em; font-size: 1em; animation: slide-in 0.4s ease-out backwards; }
  @keyframes slide-in { from { opacity: 0; transform: translateY(12px); } }
  .board-row.r1 { background: #FFE08A; }
  .board-row.r2 { background: #E6E6E6; }
  .board-row.r3 { background: #F3C79B; }
  .board-row.mine { background: var(--ink); color: var(--cream); }
  .board-rank { width: 2.2em; flex: none; text-align: center; font-family: 'Black Han Sans', sans-serif; }
  .board-name { flex: 1; min-width: 0; word-break: break-all; }
  .board-extra { font-size: 0.75em; opacity: 0.85; }
  .board-score { flex: none; font-family: 'Black Han Sans', sans-serif; min-width: 4.5em; text-align: right; }
  .board-gap { flex: none; text-align: center; line-height: 0.8; opacity: 0.5; }
  .board-more { flex: none; text-align: center; font-size: 0.7em; opacity: 0.7; }
  .move { flex: none; font-size: 0.7em; min-width: 2.6em; text-align: center; }
  .move.up { color: #26890C; }
  .move.down { color: #E21B3C; }
  .board-row.mine .move.up { color: #7CE06A; }
  .board-row.mine .move.down { color: #FF8A8A; }

  /* ===== 최종 결과: 시상대 ===== */
  .final-stage { flex: 1 1 auto !important; min-height: 0; width: 100%; position: relative;
                 display: flex; flex-direction: column; align-items: center; justify-content: center; gap: 1vmin; }
  .final-title { font-family: 'Black Han Sans', sans-serif; color: var(--cream); text-align: center;
                 font-size: min(9vmin, 120px); line-height: 1.05; }
  .final-sub { font-size: clamp(16px, 3vmin, 28px); color: var(--cream); text-align: center; }
  .podium { display: flex; align-items: flex-end; justify-content: center; gap: 1.5vmin; width: 100%; max-width: 900px; }
  .pod-col { flex: 1; max-width: 260px; min-width: 0; display: flex; flex-direction: column; align-items: center;
             animation: rise 0.7s ease-out both; }
  @keyframes rise { from { opacity: 0; transform: translateY(60px); } }
  .pod-name { font-size: clamp(13px, 2.8vmin, 34px); text-align: center; word-break: break-all; max-width: 100%;
              background: var(--cream); border-radius: 10px; padding: 2px 0.6em; margin-bottom: 0.4vmin; }
  .pod-score { font-size: clamp(12px, 2vmin, 22px); margin-bottom: 0.6vmin; }
  .pod-block { width: 100%; border-radius: 1.4vmin 1.4vmin 0 0; display: flex; flex-direction: column; align-items: center;
               padding-top: 1vmin; font-family: 'Black Han Sans', sans-serif; font-size: min(6vmin, 72px); color: white;
               box-shadow: inset 0 -8px 0 rgba(0,0,0,0.12); }
  .p1 .pod-block { height: min(24dvh, 240px); background: #F2B705; }
  .p2 .pod-block { height: min(16dvh, 170px); background: #A9AFB8; }
  .p3 .pod-block { height: min(11dvh, 120px); background: #C77D3A; }
  .pod-char { position: relative; display: flex; justify-content: center; }
  .pod-char img { height: min(18dvh, 180px); width: auto; animation: bob 2.4s ease-in-out infinite; }
  .crown { position: absolute; top: -3.5vmin; font-size: min(6dvh, 64px); animation: wobble 1.6s ease-in-out infinite; z-index: 1; }
  @keyframes wobble { 0%, 100% { transform: rotate(-10deg); } 50% { transform: rotate(10deg); } }
  .podium.sm .p1 .pod-block { height: min(16dvh, 160px); }
  .podium.sm .p2 .pod-block { height: min(11dvh, 110px); }
  .podium.sm .p3 .pod-block { height: min(7dvh, 70px); }
  .podium.sm .pod-char img { height: min(12dvh, 120px); }
  /* ===== 시상식 (뒤에 큰 문구, 앞에 단상과 용강이들) ===== */
  .ceremony { --fig: min(25dvh, 25vw); position: relative; width: 100%; max-width: 1400px; flex: 1 1 auto;
              min-height: 0; background: var(--cream); border-radius: 2.4vmin; overflow: hidden; }
  .bg-phrase { position: absolute; left: 3%; right: 3%; top: 2%; height: 60%; overflow: hidden;
               display: flex; flex-direction: column; align-items: center; justify-content: center; text-align: center;
               font-family: 'Black Han Sans', sans-serif; font-size: min(14vw, 18dvh); line-height: 0.98;
               letter-spacing: -0.01em; white-space: nowrap; animation: phrase-in 0.9s ease-out both; }
  @keyframes phrase-in { from { opacity: 0; transform: scale(1.06); } }
  .podium-zone { position: absolute; left: 0; right: 0; bottom: 0; display: flex; justify-content: center;
                 align-items: flex-end; gap: 1.2vmin; padding: 0 2vmin; }
  .podium-zone.has-me { padding-right: calc(var(--fig) * 0.75 + 4vmin); }
  .pod-col { width: calc(var(--fig) * 0.95); display: flex; flex-direction: column; align-items: center;
             animation: rise 0.7s ease-out both; }
  @keyframes rise { from { opacity: 0; transform: translateY(60px); } }
  /* 용강이 머리 위 이름표: 이름(또는 팀 이름)과 점수를 크게 */
  .pod-tag { width: max-content; min-width: 100%; max-width: calc(100% + 6vmin); box-sizing: border-box;
             background: white; border-radius: 1.6vmin; padding: 0.3em 0.7em 0.3em; margin-bottom: 1.8vmin;
             display: flex; flex-direction: column; align-items: center; text-align: center; line-height: 1.12;
             font-size: clamp(15px, 3.6vmin, 46px); box-shadow: 0 0.5vmin 1.2vmin rgba(0,0,0,0.18);
             border-bottom: 0.9vmin solid var(--tc, #D9D2CB); }
  .pod-tag b { max-width: 100%; font-weight: normal; font-family: 'Black Han Sans', sans-serif; word-break: break-all;
               display: -webkit-box; -webkit-line-clamp: 2; -webkit-box-orient: vertical; overflow: hidden;
               text-wrap: balance; }
  .pod-tag b small { font-size: 0.5em; opacity: 0.8; }
  .pod-tag .sc { font-family: 'Black Han Sans', sans-serif; font-size: 1.12em; white-space: nowrap; }
  .pod-tag .rk { font-size: 0.6em; opacity: 0.85; margin-top: 0.1em; }
  .pod-tag .sub { font-style: normal; font-size: 0.5em; margin-top: 0.2em; opacity: 0.8; }
  .pod-tag .sub.chip { opacity: 1; color: white; background: var(--tc); border-radius: 999px; padding: 0.1em 0.8em;
                       text-shadow: 0 1px 2px rgba(0,0,0,0.4); }
  .pod-tag.me { background: var(--ink); color: var(--cream); }
  .pod-fig { position: relative; height: var(--fig); }
  .p1 .pod-fig { height: calc(var(--fig) * 1.1); }
  .pod-fig img { height: 100%; width: auto; display: block; filter: drop-shadow(0 0.6vmin 0.6vmin rgba(0,0,0,0.2));
                 animation: hop 1.1s ease-in-out infinite; transform-origin: bottom center; }
  .p2 .pod-fig img { animation-delay: 0.2s; }
  .p3 .pod-fig img { animation-delay: 0.4s; }
  @keyframes hop { 0%, 100% { transform: translateY(0) scaleY(1); } 15% { transform: translateY(0) scaleY(0.94); }
                   45% { transform: translateY(-7%) scaleY(1.02); } }
  .prize { position: absolute; left: 50%; bottom: 16%; transform: translateX(-50%); line-height: 1;
           font-size: calc(var(--fig) * 0.24); animation: cheer 1.1s ease-in-out infinite;
           filter: drop-shadow(0 0.4vmin 0.4vmin rgba(0,0,0,0.3)); }
  .p1 .prize { font-size: calc(var(--fig) * 0.36); bottom: 12%; }
  @keyframes cheer { 0%, 100% { transform: translate(-50%, 0) rotate(-6deg); }
                     45% { transform: translate(-50%, -22%) rotate(6deg); } }
  .pod-block { width: 100%; border-radius: 1.4vmin 1.4vmin 0 0; display: flex; justify-content: center;
               padding-top: 0.6vmin; font-family: 'Black Han Sans', sans-serif; font-size: calc(var(--fig) * 0.22);
               color: white; box-shadow: inset 0 -0.8vmin 0 rgba(0,0,0,0.12); }
  .p1 .pod-block { height: calc(var(--fig) * 0.72); background: #F2B705; }
  .p2 .pod-block { height: calc(var(--fig) * 0.48); background: #A9AFB8; }
  .p3 .pod-block { height: calc(var(--fig) * 0.30); background: #C77D3A; }
  .envy-col { position: absolute; right: 2.5%; bottom: 0; width: calc(var(--fig) * 0.75); display: flex;
              flex-direction: column; align-items: center; animation: rise 0.7s ease-out 3.4s both; }
  .envy-col .pod-fig { height: calc(var(--fig) * 0.9); }
  .envy-col .pod-fig img { animation: sway 2.4s ease-in-out infinite; }
  @keyframes sway { 0%, 100% { transform: rotate(-2deg); } 50% { transform: rotate(2deg) translateX(-3%); } }
  .sigh { position: absolute; top: 4%; right: 2%; font-size: calc(var(--fig) * 0.14); animation: drip 2s ease-in infinite; }
  @keyframes drip { 0% { opacity: 0; transform: translateY(-20%); } 30% { opacity: 1; } 100% { opacity: 0; transform: translateY(60%); } }
  /* 세로 화면(휴대폰): 폭이 좁으니 1·2·3등 단상을 화면 폭에 맞춰 3등분하고,
     3등 밖인 내 용강이는 단상 아래 바닥에 서서 올려다봐요 */
  @media (orientation: portrait) {
    .ceremony { --fig: min(24dvh, 33vw); }
    .pod-tag { font-size: clamp(14px, 4.2vw, 24px); }
    .pod-col { flex: 1 1 0; width: auto; min-width: 0; max-width: 36%; }
    .podium-zone { gap: 1vmin; padding: 0 1vmin; }
    .podium-zone.has-me { padding-right: 1vmin; bottom: calc(var(--fig) * 0.8 + 22vmin); }
    .envy-col { right: 3%; width: calc(var(--fig) * 1.15); }
    .envy-col .pod-fig { height: calc(var(--fig) * 0.8); }
    .bg-phrase { top: 18%; height: 46%; font-size: min(15vw, 12dvh); }
    .p1 .pod-block { height: calc(var(--fig) * 0.72 + 8dvh); }
    .p2 .pod-block { height: calc(var(--fig) * 0.48 + 5dvh); }
    .p3 .pod-block { height: calc(var(--fig) * 0.30 + 3dvh); }
    .podium-zone.has-me .p1 .pod-block { height: calc(var(--fig) * 0.72 + 3dvh); }
    .podium-zone.has-me .p2 .pod-block { height: calc(var(--fig) * 0.48 + 2dvh); }
    .podium-zone.has-me .p3 .pod-block { height: calc(var(--fig) * 0.30 + 1dvh); }
  }
  .ceremony-top { position: absolute; top: 2%; left: 0; right: 0; display: flex; justify-content: center; z-index: 3; }
  .confetti { position: fixed; inset: 0; pointer-events: none; overflow: hidden; z-index: 40; }
  .confetti span { position: absolute; top: -20px; width: 10px; height: 16px; border-radius: 2px; opacity: 0;
                   animation: fall linear both; }
  @keyframes fall { 0% { opacity: 1; transform: translateY(0) rotate(0); }
                    100% { opacity: 1; transform: translateY(110vh) rotate(720deg); } }
  .team-win { font-family: 'Black Han Sans', sans-serif; color: white; border-radius: 1.6vmin; text-align: center;
              padding: 0.2em 1em; font-size: min(4vmin, 44px); animation: rise 0.7s ease-out both; animation-delay: 3.4s; }

  /* ===== 소리 버튼 ===== */
  .sound-btn { position: fixed; right: 18px; bottom: calc(18px + env(safe-area-inset-bottom, 0px));
               z-index: 50; background: var(--ink); color: var(--cream); border: none; cursor: pointer;
               border-radius: 999px; padding: 12px 22px; font-family: 'Jua', sans-serif; font-size: 18px;
               box-shadow: 0 4px 0 rgba(0,0,0,0.25); }
  .sound-btn:active { transform: translateY(2px); box-shadow: 0 2px 0 rgba(0,0,0,0.25); }
</style>
<script>
  // 글자가 자기 칸보다 크면 들어갈 때까지 조금씩 줄이기 (class="autofit" 인 것들)
  window.fitText = function () {
    document.querySelectorAll('.autofit').forEach(el => {
      el.style.fontSize = '';
      let size = parseFloat(getComputedStyle(el).fontSize), guard = 0;
      while ((el.scrollHeight > el.clientHeight + 1 || el.scrollWidth > el.clientWidth + 1)
             && size > 11 && guard++ < 80) {
        size -= 1;
        el.style.fontSize = size + 'px';
      }
    });
  };
  (function () {
    let timer = null;
    const later = () => { clearTimeout(timer); timer = setTimeout(window.fitText, 60); };
    window.addEventListener('resize', later);
    window.addEventListener('orientationchange', later);
    new MutationObserver(later).observe(document.documentElement, { childList: true, subtree: true });
    if (document.fonts) document.fonts.ready.then(later);
  })();
</script>
""", shared=True)

BITS = "".join(
    f'<span class="bit" style="left:{x}%; top:{y}%; width:{w}px; height:{int(w * 0.55)}px; '
    f'background:{c}; animation-delay:{d}s"></span>'
    for x, y, w, c, d in [
        (4, 18, 56, "#F5A33A", 0), (14, 62, 40, "#1E5BB8", 1.2), (24, 10, 30, "#F4EFEA", 2.1),
        (8, 85, 46, "#F5A33A", 0.6), (70, 8, 50, "#F5A33A", 1.6), (88, 22, 36, "#1E5BB8", 0.3),
        (93, 58, 54, "#F5A33A", 2.4), (78, 80, 32, "#F4EFEA", 1.0), (62, 48, 26, "#1E5BB8", 2.8),
        (33, 74, 34, "#F5A33A", 1.9),
    ]
)


# =====================================================================
#  문제 화면 그리기 (진행자 화면과 참가자 화면이 같이 써요)
# =====================================================================
def btn(text, color, on_click, text_color="white"):
    """색을 직접 정한 버튼 (글씨 색도 같이 정해서 배경에 묻히지 않게)"""
    return ui.button(text, color=color, on_click=on_click).props(f"text-color={text_color}")


def text_size_class(texts):
    """보기 글자 수에 따라 크기 단계 정하기 (짧을수록 크게).
    한 문제의 카드 4장은 가장 긴 보기에 맞춰 같은 크기로 보여줘요 (들쭉날쭉하지 않게)"""
    n = max(len(str(t).replace(" ", "")) for t in texts)
    if n <= 3:
        return " len-s"
    if n <= 8:
        return " len-m"
    if n <= 16:
        return " len-l"
    return " len-xl"


def card_look(i, is_ox):
    """i번째 카드의 (추가 class, style). 그림이 있으면 그림 카드, 없으면 색깔 카드"""
    color = (OX_COLORS if is_ox else COLORS)[i]
    files = OX_CARD_IMAGES if is_ox else CARD_IMAGES
    style = f"background: {color}; --c: {color}; --i: {i};"
    if i < len(files) and (CARD_DIR / files[i]).exists():
        return " has-img", style + f" --bg: url('/images/cards/{files[i]}');"
    return "", style


def draw_question_stage(q, live, on_pick=None, on_text=None, host=False):
    """위에는 큰 문제, 아래에는 남은 공간을 꽉 채우는 카드 보기.
    on_pick: 카드를 눌렀을 때 할 일 (진행자 화면은 None이라 눌러도 아무 일 없음)"""
    with ui.element("div").classes("quiz-stage"):
        ui.label(q["문제"]).classes("stage-title autofit")
        ui.label(f"{game['번호'] + 1}번 문제 / 전체 {len(questions)}문제  ·  {q['유형']}").classes("stage-sub")
        with ui.row().classes("items-center gap-2 no-wrap"):
            live["시간"] = ui.label().classes("stage-pill")
            if host:
                live["응답"] = ui.label().classes("stage-pill")
        if q["사진"]:
            ui.html(photo_html(q, "q-photo"), sanitize=False).classes("contents")

        if q["유형"] in ("4지선다", "OX"):
            is_ox = q["유형"] == "OX"
            texts = ["⭕", "❌"] if is_ox else q["보기"]
            colors = OX_COLORS if is_ox else COLORS
            with ui.element("div").classes("choice-grid two" if is_ox else "choice-grid"):
                for i, text in enumerate(texts):
                    look, style = card_look(i, is_ox)
                    card = ui.element("div").classes(
                        ("choice-card enter" if on_pick else "choice-card static enter") + look
                    ).style(style).mark(f"choice-{i + 1}")
                    with card:
                        if look:
                            ui.element("span").classes("shine")     # 반짝 지나가는 빛
                        if not is_ox:
                            with ui.element("div").classes("card-top"):
                                ui.label(str(i + 1)).classes("choice-num")
                        ui.label(text).classes("choice-text ox" if is_ox else "choice-text autofit" + text_size_class(texts))
                    if on_pick:
                        card.on("click", lambda e, n=i + 1: on_pick(n))
        else:  # 단답형
            if on_text:
                with ui.element("div").classes("short-box"):
                    answer_box = ui.input("정답 입력").props("outlined")
                    answer_box.on("keydown.enter", lambda: on_text(answer_box.value))
                    btn("제출", color="#2B2523",
                              on_click=lambda: on_text(answer_box.value)).classes("text-xl")
            else:
                ui.label("✍️ 휴대폰에 정답을 직접 입력하세요").classes("msg").style("margin: auto 0")


def draw_answer_stage(q, title, sub, pills, my_answer=None):
    """정답 공개: 문제 화면과 같은 카드 위에 고른 인원이 차오르고, 정답 카드만 밝게.
    my_answer: 참가자 화면이면 그 사람이 고른 번호 ('내 선택' 표시용)"""
    ui.html('<div class="flash"></div>', sanitize=False).classes("contents")   # 번쩍!
    with ui.element("div").classes("quiz-stage reveal-stage"):
        ui.label(title).classes("stage-title autofit")
        ui.label(sub).classes("stage-sub")
        with ui.row().classes("items-center justify-center gap-2"):
            for text in pills:
                ui.label(text).classes("stage-pill")

        if q["유형"] not in ("4지선다", "OX"):      # 단답형은 막대그래프로
            ui.html(bar_chart_html(q), sanitize=False).classes("contents")
            return

        is_ox = q["유형"] == "OX"
        texts = ["⭕", "❌"] if is_ox else q["보기"]
        colors = OX_COLORS if is_ox else COLORS
        counts = [sum(1 for p in players.values() if p["답"] == n) for n in range(1, len(texts) + 1)]
        biggest = max(counts) or 1
        with ui.element("div").classes("choice-grid two" if is_ox else "choice-grid"):
            for i, text in enumerate(texts):
                right = (i + 1) == q["정답"]
                look, style = card_look(i, is_ox)
                with ui.element("div").classes(
                    "choice-card static reveal " + ("right" if right else "wrong") + look
                ).style(style):
                    height = int(100 * counts[i] / biggest)
                    ui.html(f'<div class="fill" style="height:{height}%"></div>', sanitize=False).classes("contents")
                    with ui.element("div").classes("card-top on-top"):
                        badge = "✅" if right else ("" if is_ox else str(i + 1))
                        ui.label(badge).classes("choice-num" if badge else "")
                        ui.label(f"{counts[i]}명").classes("choice-count")
                    ui.label(text).classes("choice-text ox" if is_ox else "choice-text autofit" + text_size_class(texts))
                    if my_answer == i + 1:
                        ui.label("👆 내 선택").classes("my-pick")


def draw_board(kind, my_name=None, limit=8):
    """점수판. kind: "라운드"(이번 라운드 점수) / "누적"(누적 순위) / "최종"
    상위 limit명만 보여주고, 내가 그 밖이면 맨 아래에 내 줄을 따로 보여줘요."""
    with ui.element("div").classes("quiz-stage"):
        before_rank = {}
        if kind == "라운드":
            q = questions[game["번호"]]
            ui.label("이번 라운드 점수").classes("stage-title autofit")
            ui.label(f"{game['번호'] + 1}번 문제  ·  정답: {answer_text(q)}").classes("stage-sub")
            rows = sorted(players.items(), key=lambda kv: kv[1]["이번점수"], reverse=True)
        elif kind == "최종":
            ui.label("전체 순위").classes("stage-title autofit")
            ui.label(f"전체 {len(questions)}문제 최종 점수").classes("stage-sub")
            rows = ranking()
        else:
            ui.label("🏆 누적 순위").classes("stage-title autofit")
            ui.label(f"{game['번호'] + 1}번 문제까지  ·  전체 {len(questions)}문제").classes("stage-sub")
            rows = ranking()
            # 이번 문제 전의 순위 (순위가 몇 칸 올랐는지 보여주려고)
            before = sorted(players.items(), key=lambda kv: kv[1]["점수"] - kv[1]["이번점수"], reverse=True)
            before_rank = {name: i for i, (name, _) in enumerate(before, 1)}

        if not rows:
            ui.label("참가자가 없어요").classes("stage-sub")
            return

        medals = {1: "🥇", 2: "🥈", 3: "🥉"}

        def row(rank, name, p):
            css = "board-row"
            if kind != "라운드" and rank <= 3:
                css += f" r{rank}"
            if name == my_name:
                css += " mine"
            with ui.element("div").classes(css).style(f"animation-delay: {min(rank * 0.06, 0.8)}s"):
                ui.label(str(rank) if kind == "라운드" else medals.get(rank, str(rank))).classes("board-rank")
                with ui.column().classes("board-name gap-0"):
                    ui.label(name + ("  (나)" if name == my_name else ""))
                    if kind == "라운드" and p["보너스"]:
                        ui.label(f"🔥 {p['연속']}연속 보너스 +{p['보너스']}").classes("board-extra")
                    if kind == "누적" and p["연속"] >= STREAK_FROM:
                        ui.label(f"🔥{p['연속']}연속 정답 중").classes("board-extra")
                    if kind == "최종":
                        right = sum(1 for r in p["기록"] if r["맞음"])
                        ui.label(f"맞힌 문제 {right} / {len(questions)}").classes("board-extra")
                if kind == "라운드":
                    if p["이번점수"] > 0:
                        ui.label(f"+{p['이번점수']}").classes("board-score")
                    elif p["답"] is None:
                        ui.label("⏰ 0").classes("board-score")
                    else:
                        ui.label("❌ 0").classes("board-score")
                else:
                    if kind == "누적":
                        moved = before_rank[name] - rank
                        if moved > 0:
                            ui.label(f"▲{moved}").classes("move up")
                        elif moved < 0:
                            ui.label(f"▼{-moved}").classes("move down")
                        else:
                            ui.label("–").classes("move")
                    ui.label(f"{p['점수']}점").classes("board-score")

        names = [n for n, _ in rows]
        my_rank = names.index(my_name) + 1 if my_name in names else None
        shown = rows[:limit]
        with ui.element("div").classes("board autofit"):
            for rank, (name, p) in enumerate(shown, 1):
                row(rank, name, p)
            if my_rank and my_rank > limit:
                ui.label("⋮").classes("board-gap")
                row(my_rank, my_name, players[my_name])
            hidden = len(rows) - len(shown) - (1 if my_rank and my_rank > limit else 0)
            if hidden > 0:
                ui.label(f"외 {hidden}명").classes("board-more")


def ceremony_html(kind="개인", my_name=None, my_team=None):
    """최종 시상식: 뒤에는 큰 문구, 앞에는 단상 위 용강이들.
    kind="개인": 닉네임 + 점수 (+ 소속 팀),  kind="팀": 팀 이름 + 팀 점수
    1등은 트로피, 2·3등은 은·동메달을 들고 기뻐해요.
    참가자 화면에서 내가(내 팀이) 3등 안에 없으면, 단상 옆에서 부러운 얼굴로 바라보는 용강이가 나와요."""
    happy = "yg_happy.png" if (IMAGE_DIR / "yg_happy.png").exists() else "yonggang.png"
    envy = "yg_envy.png" if (IMAGE_DIR / "yg_envy.png").exists() else "yonggang.png"
    delays = {1: 2.6, 2: 1.3, 3: 0.4}        # 몇 초 뒤에 나타날지 (1등이 마지막!)
    prizes = {1: "🏆", 2: "🥈", 3: "🥉"}

    # 순위표를 같은 모양의 목록으로 만들기: 이름표에 쓸 글자들
    if kind == "팀":
        entries = [{"key": t, "title": t, "score": sc, "color": team_color(t),
                    "sub": f"팀원 {cnt}명", "chip": False}
                   for t, sc, cnt in team_rows()]
        my_key, mark, phrase_lines, who = my_team, " (우리 팀)", FINAL_PHRASE_TEAM, "우리 팀 "
    else:
        entries = [{"key": n, "title": n, "score": p["점수"],
                    "color": team_color(p["팀"]) if p["팀"] in teams else "#D9D2CB",
                    "sub": p["팀"] if p["팀"] in teams else "", "chip": True}
                   for n, p in ranking()]
        my_key, mark, phrase_lines, who = my_name, " (나)", FINAL_PHRASE, ""
    top = entries[:3]
    keys = [e["key"] for e in entries]
    my_rank = keys.index(my_key) + 1 if my_key in keys else None

    def tag(e, rank_text=""):
        me = e["key"] == my_key
        sub = ""
        if e["sub"]:
            sub = f'<i class="sub{" chip" if e["chip"] else ""}">{html.escape(e["sub"])}</i>'
        return (f'<div class="pod-tag{" me" if me else ""}" style="--tc:{e["color"]}">'
                f'<b>{html.escape(e["title"])}{f"<small>{mark}</small>" if me else ""}</b>'
                f'{f"<span class=rk>{rank_text}</span>" if rank_text else ""}'
                f'<span class="sc">{e["score"]}점</span>{sub}</div>')

    cols = []
    for rank in (2, 1, 3):                     # 왼쪽부터 2등, 1등, 3등
        if rank > len(top):
            continue
        e = top[rank - 1]
        cols.append(
            f'<div class="pod-col p{rank}" style="animation-delay:{delays[rank]}s">'
            f'{tag(e)}'
            f'<div class="pod-fig"><img src="/images/{happy}" alt="">'
            f'<span class="prize">{prizes[rank]}</span></div>'
            f'<div class="pod-block">{rank}</div></div>')

    me_html = ""
    if my_rank and my_rank > 3:
        e = entries[my_rank - 1]
        me_html = (f'<div class="envy-col">{tag(e, f"{who}{my_rank}등")}'
                   f'<div class="pod-fig"><img src="/images/{envy}" alt="">'
                   f'<span class="sigh">💧</span></div></div>')

    phrase = "<br>".join(html.escape(line) for line in phrase_lines)
    rnd = random.Random(len(players))          # 꽃가루 위치
    pieces = "".join(
        f'<span style="left:{rnd.uniform(0, 100):.1f}%; background:{rnd.choice(CONFETTI_COLORS)}; '
        f'animation-delay:{2.6 + rnd.uniform(0, 2.5):.2f}s; animation-duration:{rnd.uniform(2.5, 4.5):.2f}s; '
        f'transform: rotate({rnd.randint(0, 90)}deg)"></span>'
        for _ in range(70))
    zone = "podium-zone has-me" if me_html else "podium-zone"
    return (f'<div class="confetti">{pieces}</div>'
            f'<div class="ceremony">'
            f'<div class="bg-phrase autofit" style="color:{FINAL_PHRASE_COLOR}">{phrase}</div>'
            f'<div class="{zone}">{"".join(cols)}</div>{me_html}</div>')


def draw_team_board(kind, my_team=None, limit=8):
    """팀 점수판. kind: "라운드" / "누적" / "최종" """
    if not teams:
        return
    rows = team_rows("이번점수" if kind == "라운드" else "점수")
    titles = {"라운드": "👥 팀별 이번 라운드", "누적": "👥 팀 순위", "최종": "👥 최종 팀 순위"}
    medals = {1: "🥇", 2: "🥈", 3: "🥉"}
    with ui.element("div").classes("quiz-stage"):
        ui.label(titles[kind]).classes("stage-title autofit")
        ui.label(f"팀원 점수 {TEAM_SCORE}으로 계산" if TEAM_SCORE == "평균"
                 else f"팀원 점수 {TEAM_SCORE}로 계산").classes("stage-sub")
        if not rows:
            ui.label("팀에 들어간 사람이 없어요").classes("stage-sub")
            return
        with ui.element("div").classes("board autofit"):
            for rank, (t, score, count) in enumerate(rows[:limit], 1):
                css = "board-row"
                if kind != "라운드" and rank <= 3:
                    css += f" r{rank}"
                if t == my_team:
                    css += " mine"
                with ui.element("div").classes(css).style(
                        f"animation-delay: {min(rank * 0.08, 0.8)}s; border-left: 10px solid {team_color(t)}"):
                    ui.label(str(rank) if kind == "라운드" else medals.get(rank, str(rank))).classes("board-rank")
                    with ui.column().classes("board-name gap-0"):
                        ui.label(t + ("  (우리 팀)" if t == my_team else ""))
                        ui.label(f"{count}명").classes("board-extra")
                    ui.label(f"+{score}" if kind == "라운드" else f"{score}점").classes("board-score")
            if len(rows) > limit:
                ui.label(f"외 {len(rows) - limit}팀").classes("board-more")


def team_picker(on_pick, selected=None):
    """팀 고르기 버튼들 (고른 팀만 색이 꽉 차요)"""
    buttons = {}

    def pick(t):
        for name, b in buttons.items():
            if name == t:
                b.props(remove="outline")
            else:
                b.props("outline")
        on_pick(t)

    with ui.element("div").classes("team-pick"):
        for t in teams:
            b = ui.button(t, color=team_color(t), on_click=lambda t=t: pick(t))
            if t != selected:
                b.props("outline")
            buttons[t] = b


# =====================================================================
#  소리 (진행자 화면에서만 나와요)
#  - 배경음악: sounds 폴더의 mp3 파일
#  - 효과음: 파일 없이 브라우저가 직접 만들어요 (띡띡, 딩동댕, 빰빠밤)
# =====================================================================
SOUND_SCRIPT = """
<script>
window.quizSound = (() => {
  let enabled = false, scene = null, ctx = null;
  const MUSIC = { '대기': 'lobby', '순위': 'lobby', '누적': 'lobby', '끝': 'lobby', '팀끝': 'lobby', '최종표': 'lobby', '문제': 'quiz', '정답': null };
  const VOLUME = %VOLUME%;

  function playMusic(which, restart) {
    for (const [name, id] of [['lobby', 'bgm-lobby'], ['quiz', 'bgm-quiz']]) {
      const el = document.getElementById(id);
      if (!el) continue;
      if (name === which) {
        if (restart) el.currentTime = 0;
        el.volume = VOLUME;
        if (el.paused) el.play().catch(() => {});
      } else {
        fadeOut(el);
      }
    }
  }

  // 음악을 뚝 끊지 않고 0.35초 동안 부드럽게 줄이면서 끄기
  function fadeOut(el, ms = 350) {
    if (!el || el.paused || el.dataset.fading) return;
    el.dataset.fading = '1';
    const start = el.volume, t0 = performance.now();
    const step = (now) => {
      const k = Math.min(1, (now - t0) / ms);
      el.volume = start * (1 - k);
      if (k < 1) requestAnimationFrame(step);
      else { el.pause(); el.volume = VOLUME; delete el.dataset.fading; }
    };
    requestAnimationFrame(step);
  }

  // 정답 공개 효과음 (파일이 없으면 딩동댕으로 대신)
  function playReveal() {
    const el = document.getElementById('sfx-reveal');
    if (!el) { effects.reveal(); return; }
    el.currentTime = 0;
    el.volume = %REVEAL_VOLUME%;
    el.play().catch(() => effects.reveal());
  }

  function tone(freq, start, dur, type = 'sine', vol = 0.25) {
    if (!enabled) return;
    ctx = ctx || new (window.AudioContext || window.webkitAudioContext)();
    const osc = ctx.createOscillator(), gain = ctx.createGain();
    const t = ctx.currentTime + start;
    osc.type = type;
    osc.frequency.value = freq;
    gain.gain.setValueAtTime(vol, t);
    gain.gain.exponentialRampToValueAtTime(0.001, t + dur);
    osc.connect(gain); gain.connect(ctx.destination);
    osc.start(t); osc.stop(t + dur);
  }

  const effects = {
    start:   () => { tone(392, 0, 0.15, 'triangle'); tone(523, 0.15, 0.3, 'triangle'); },
    tick:    () => tone(880, 0, 0.12, 'square', 0.12),
    reveal:  () => { tone(523, 0, 0.3); tone(659, 0.2, 0.3); tone(784, 0.4, 0.6); },
    fanfare: (at = 0) => { [523, 659, 784].forEach((f, i) => tone(f, at + i * 0.15, 0.3, 'triangle'));
                           tone(1047, at + 0.45, 0.9, 'triangle'); },
    final:   () => { for (let i = 0; i < 24; i++) tone(150 + (i % 2) * 20, i * 0.1, 0.09, 'triangle', 0.12);
                     effects.fanfare(2.6); },   // 두구두구… 1등 발표 순간 빠바밤!
  };

  return {
    setScene(s) {
      scene = s;
      if (!enabled) return;
      playMusic(MUSIC[s], s === '문제');      // 문제마다 음악을 처음부터
      if (s === '문제') effects.start();
      if (s === '정답') playReveal();
      if (s === '끝' || s === '팀끝') effects.final();
    },
    effect(name) { if (enabled && effects[name]) effects[name](); },
    toggle(button) {
      enabled = !enabled;
      if (enabled) {
        ctx = ctx || new (window.AudioContext || window.webkitAudioContext)();
        ctx.resume();
        playMusic(MUSIC[scene], false);
      } else {
        playMusic(null);
      }
      button.textContent = enabled ? '🔊 소리 끄기' : '🔇 소리 켜기';
    },
  };
})();
</script>
"""


def install_sound():
    """진행자 화면에 음악 플레이어, 소리 버튼 붙이기"""
    ui.add_head_html(SOUND_SCRIPT.replace("%VOLUME%", str(MUSIC_VOLUME))
                     .replace("%REVEAL_VOLUME%", str(REVEAL_VOLUME)))
    reveal = ""
    if (SOUND_DIR / REVEAL_SOUND).exists():
        reveal = f'<audio id="sfx-reveal" src="/sounds/{REVEAL_SOUND}" preload="auto"></audio>'
    ui.html(
        reveal +
        f'<audio id="bgm-lobby" src="/sounds/{LOBBY_MUSIC}" loop preload="auto"></audio>'
        f'<audio id="bgm-quiz" src="/sounds/{QUIZ_MUSIC}" loop preload="auto"></audio>'
        '<button class="sound-btn" onclick="quizSound.toggle(this)">🔇 소리 켜기</button>',
        sanitize=False)


# =====================================================================
#  진행자 비밀번호
# =====================================================================
# =====================================================================
#  입장용 QR코드
# =====================================================================
def lan_ip():
    """이 컴퓨터의 와이파이 주소 (예: 172.16.1.215)"""
    import socket
    sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        sock.connect(("10.255.255.255", 1))
        return sock.getsockname()[0]
    except Exception:
        return None
    finally:
        sock.close()


def join_url_for(origin):
    """참가자가 들어올 주소 정하기.
    - 진행자가 통로 주소(https://…trycloudflare.com/host)로 열었으면 그 주소 그대로
    - localhost 로 열었으면, 휴대폰은 localhost 로 못 들어오니 와이파이 주소로 바꿔요"""
    if JOIN_URL:
        return JOIN_URL
    m = re.match(r"^(https?)://([^/:]+)(?::(\d+))?", origin or "")
    if m and m.group(2) not in ("localhost", "127.0.0.1", "0.0.0.0", "[::1]"):
        return origin.rstrip("/") + "/"
    ip = lan_ip()
    port = (m.group(3) if m else None) or "8080"
    return f"http://{ip or 'localhost'}:{port}/"


def qr_svg(url):
    if segno is None:
        return ""
    return segno.make(url, error="m").svg_inline(scale=1, border=2, dark="#2B2523", light="#FFFFFF",
                                                 omitsize=True)


def qr_board_html(url):
    """나무 안내 팻말 속 QR (글자 없이 QR만)"""
    svg = qr_svg(url)
    if not svg:
        svg = '<div class="qm-msg">QR을 보려면<br>명령 프롬프트에서<br><b>pip install segno</b></div>'
    return f'<div class="qm-board">{svg}</div>'


# 팻말 그림 속 판자(QR 자리) 위치와, 그림 속 용강이를 원래 용강이 자리에 맞추는 값 (그림을 만들 때 잰 값)
SIGN_IMAGES = {
    "보통": {"file": "yg_sign.png", "w": 818, "h": 1214, "left": 65.08, "top": 23.56, "size": 43.13, "angle": 19.13,
             "slot": {"height": 167.23, "top": -56.04, "left": 1.55, "ox": 29.1, "oy": 90.4}},
    "크게": {"file": "yg_sign2.png", "w": 1154, "h": 1407, "left": 62.3, "top": 30.25,
             "size": 45.86, "angle": 19.13,
             "slot": {"height": 192.95, "top": -81.84, "left": -2.8,
                      "ox": 22.6, "oy": 91.7}},
}


def sign_spec():
    """설정한 크기의 팻말 그림 (파일이 없으면 다른 크기로)"""
    for key in (SIGN_SIZE, "크게", "보통"):
        sp = SIGN_IMAGES.get(key)
        if sp and (IMAGE_DIR / sp["file"]).exists():
            return sp
    return None


def qr_mascot_html(url):
    """용강이가 팻말을 번쩍 들고 낑낑대는 모습 + 판자 위 QR (누르면 크게)"""
    opener = "document.getElementById('qr-big').classList.add('show')"
    sp = sign_spec()
    if sp:
        svg = qr_svg(url) or '<div class="qm-msg">pip install segno</div>'
        sl = sp["slot"]
        return (f'<div class="qr-sign" onclick="{opener}" '
                f'style="aspect-ratio: {sp["w"]} / {sp["h"]}; --sh:{sl["height"]}%; --st:{sl["top"]}%; '
                f'--sl:{sl["left"]}%; transform-origin:{sl["ox"]}% {sl["oy"]}%">'
                f'<img src="/images/{sp["file"]}" alt="">'
                f'<div class="qs-qr" style="left:{sp["left"]}%; top:{sp["top"]}%; width:{sp["size"]}%; '
                f'--a:{sp["angle"]}deg">{svg}</div>'
                f'<span class="qs-sweat">💦</span></div>')
    # (그림 파일이 없을 때) 예전 모양: 팻말을 앞으로 들고 있는 용강이
    face = "yg_strain.png" if (IMAGE_DIR / "yg_strain.png").exists() else "yonggang.png"
    return (f'<div class="qr-mascot" onclick="document.getElementById(\'qr-big\').classList.add(\'show\')">'
            f'<img class="qm-char" src="/images/{face}" alt="">'
            f'<span class="qm-sweat s1">💦</span><span class="qm-sweat s2">💧</span>'
            f'<span class="qm-foot l"></span><span class="qm-foot r"></span>'
            f'<div class="qm-sign">{qr_board_html(url)}</div>'
            f'<span class="qm-hand l"></span><span class="qm-hand r"></span></div>')


def qr_overlay_html(url):
    """화면 가득 크게 보여주는 QR 팻말 (아무 곳이나 누르면 닫힘)"""
    return (f'<div id="qr-big" class="qr-big" onclick="this.classList.remove(\'show\')">'
            f'<div class="qr-big-box">{qr_board_html(url)}</div></div>')


def need_host_login():
    if app.storage.user.get("진행자"):
        return False
    with ui.column().classes("w-full items-center gap-4 p-4"):
        ui.label("🔒 진행자 비밀번호를 입력하세요").classes("text-3xl")
        pw_box = ui.input("비밀번호", password=True)

        def check_pw():
            if pw_box.value == HOST_PASSWORD:
                app.storage.user["진행자"] = True
                ui.navigate.reload()
            else:
                ui.notify("비밀번호가 틀렸어요")

        ui.button("들어가기", on_click=check_pw)
    return True


# =====================================================================
#  진행자 화면 (큰 화면, 주소 끝에 /host)
# =====================================================================
@ui.page("/host")
def host_page():
    if need_host_login():
        return
    dialog_spot = ui.element("div")   # 팝업창 자리 (화면을 새로 그려도 팝업이 안 사라지게 따로 둬요)
    qr_spot = ui.html("", sanitize=False)   # 크게 보는 QR 자리 (사람이 들어와도 안 닫히게 따로 둬요)
    area = ui.element("div").classes("screen")
    live = {}
    shown = {"화면": None, "장면": None, "카운트": None}
    join = {"url": None}
    # QR 보기 켜짐/꺼짐 (이 진행자 브라우저에 기억해둬요. 처음엔 꺼짐)
    qr_view = {"on": bool(app.storage.user.get("QR보기", False))}

    def toggle_qr():
        qr_view["on"] = not qr_view["on"]
        app.storage.user["QR보기"] = qr_view["on"]
        shown["화면"] = None          # 대기실 다시 그리기
    install_sound()

    async def detect_join_url():
        """진행자 화면을 연 주소를 보고, 참가자가 들어올 주소(QR)를 정해요"""
        origin = None
        try:
            origin = await ui.run_javascript("window.location.origin", timeout=3)
        except Exception:
            pass
        join["url"] = join_url_for(origin)
        qr_spot.content = qr_overlay_html(join["url"])
        shown["화면"] = None           # 대기실을 QR과 함께 다시 그리기

    ui.timer(0.3, detect_join_url, once=True)

    def screen_key():
        members = tuple(sorted((n, p["팀"] or "") for n, p in players.items()))
        return (game["단계"], game["번호"], len(players), game["세트"], len(questions),
                tuple(teams), members, qr_view["on"])

    def open_team_dialog():
        dialog_spot.clear()
        with dialog_spot, ui.dialog() as dialog, ui.card().classes("min-w-[340px]"):
            ui.label("👥 팀 설정").classes("text-2xl")
            ui.label("팀을 만들면 팀전, 하나도 없으면 개인전이에요").classes("text-sm")
            team_list = ui.column().classes("w-full gap-1")

            def draw_teams():
                team_list.clear()
                with team_list:
                    if not teams:
                        ui.label("아직 팀이 없어요 (개인전)").classes("text-gray-500")
                    for t in teams:
                        with ui.row().classes("w-full items-center no-wrap"):
                            ui.element("div").classes("team-dot").style(f"background: {team_color(t)}")
                            ui.label(t).classes("text-lg flex-1")
                            ui.button(icon="close", on_click=lambda t=t: (remove_team(t), draw_teams())
                                      ).props("flat round dense size=sm")

            def add():
                if add_team(box.value):
                    box.value = ""
                    draw_teams()

            with ui.row().classes("w-full items-center no-wrap"):
                box = ui.input("새 팀 이름 (예: 1조)").classes("flex-1")
                box.on("keydown.enter", add)
                ui.button("추가", on_click=add)
            ui.button("닫기", color="grey", on_click=dialog.close).classes("self-end")
            draw_teams()
        dialog.open()

    def name_chip(n):
        with ui.element("div").classes("name-chip"):
            ui.label(n)
            ui.button(icon="close", on_click=lambda: kick(n)).props("flat round dense size=sm")

    def next_label():
        return "최종 결과 보기" if game["번호"] == len(questions) - 1 else "다음 문제 ▶"

    def draw():
        area.clear()
        live.clear()
        step = game["단계"]
        with area:
            if step == "대기":
                with ui.element("div").classes("topbar"):
                    with ui.row().classes("items-center gap-4"):
                        ui.select(list_sets(), value=game["세트"], label="퀴즈 세트",
                                  on_change=lambda e: choose_set(e.value)).classes("w-48")
                        ui.label(f"참가자 {len(players)}명  |  문제 {len(questions)}개").classes("menu")
                    ui.label("우리반 퀴즈").classes("brand")
                    with ui.row().classes("gap-2"):
                        if qr_view["on"]:      # 켜져 있으면 꽉 찬 버튼
                            btn("📱 QR 끄기", color="#1368CE", on_click=toggle_qr).classes("menu-btn")
                        else:
                            btn("📱 QR", color="#F4EFEA", on_click=toggle_qr
                                ).props("text-color=black outline").classes("menu-btn")
                        btn("👥 팀 설정", color="#F4EFEA", on_click=open_team_dialog
                                  ).props("text-color=black outline").classes("menu-btn")
                        btn("✏️ 문제 만들기", color="#F4EFEA",
                                  on_click=lambda: ui.navigate.to("/edit")
                                  ).props("text-color=black outline").classes("menu-btn")
                        btn("게임 시작!", color="#2B2523", on_click=start_game).classes("menu-btn")

                show_qr = bool(join["url"]) and qr_view["on"]
                sign_ok = sign_spec() is not None
                row_layout = show_qr and not sign_ok      # 팻말 그림이 없을 때만 예전 좌우 배치
                with ui.element("div").classes("hero has-qr" if row_layout else "hero"):
                    ui.html(BITS, sanitize=False)
                    if show_qr and sign_ok:
                        # QR을 꺼도 켜도 제목과 용강이는 같은 자리. 용강이만 팻말을 번쩍 들어요
                        ui.label("두근두근").classes("hero-title")
                        ui.label("퀴즈 타임").classes("hero-title")
                        ui.html(f'<div class="hero-sign-slot">{qr_mascot_html(join["url"])}</div>',
                                sanitize=False).classes("contents")
                    elif row_layout:
                        # 왼쪽엔 제목, 오른쪽엔 QR 팻말을 든 용강이
                        with ui.element("div").classes("hero-titles"):
                            ui.label("두근두근").classes("hero-title")
                            ui.label("퀴즈 타임").classes("hero-title")
                        ui.html(qr_mascot_html(join["url"]), sanitize=False).classes("contents")
                    else:
                        ui.label("두근두근").classes("hero-title")
                        ui.label("퀴즈 타임").classes("hero-title")
                        ui.html('<img src="/images/yonggang.png" class="hero-char">',
                                sanitize=False).classes("contents")

                with ui.element("div").classes("lobby-names"):
                    if not players:
                        ui.label("📱 QR코드를 찍고 닉네임을 입력해주세요!" if show_qr
                                 else "휴대폰으로 접속해서 닉네임을 입력해주세요!").classes("msg")
                    elif not teams:
                        with ui.element("div").classes("names"):
                            for n in players:
                                name_chip(n)
                    else:
                        # 팀전: 팀별로 묶어서 보여주기
                        with ui.element("div").classes("team-groups"):
                            for t in teams + [None]:
                                members = [n for n, p in players.items()
                                           if (p["팀"] == t if t else p["팀"] not in teams)]
                                if t is None and not members:
                                    continue
                                color = team_color(t) if t else "#8C8C8C"
                                with ui.element("div").classes("team-group").style(f"border-color: {color}"):
                                    ui.label(f"{t or '팀 미선택'}  ({len(members)}명)").classes("team-head")
                                    with ui.element("div").classes("names"):
                                        for n in members:
                                            name_chip(n)

            elif step == "문제":
                draw_question_stage(questions[game["번호"]], live, host=True)
                with ui.element("div").classes("actions"):
                    btn("바로 정답 공개", color="#2B2523", on_click=go_to_answer)

            elif step == "정답":
                q = questions[game["번호"]]
                correct = sum(1 for p in players.values() if is_correct(q, p["답"]))
                no_answer = sum(1 for p in players.values() if p["답"] is None)
                pills = [f"{len(players)}명 중 {correct}명 정답"]
                if no_answer:
                    pills.append(f"무응답 {no_answer}명")
                draw_answer_stage(q, f"정답은 {answer_text(q)}!", q["문제"], pills)
                with ui.element("div").classes("actions"):
                    btn("순위 보기", color="#2B2523", on_click=show_ranking)

            elif step in ("순위", "누적"):
                with ui.element("div").classes("boards"):
                    if step == "순위":
                        draw_board("라운드", limit=BOARD_TOP_HOST)
                        draw_team_board("라운드")
                    else:
                        draw_team_board("누적")
                        draw_board("누적", limit=BOARD_TOP_HOST)
                with ui.element("div").classes("actions"):
                    if step == "순위":
                        btn("🏆 누적 순위 보기", color="#D89E00", on_click=show_total)
                    else:
                        ui.button("← 이번 라운드 점수", color="grey", on_click=show_ranking)
                    btn(next_label(), color="#2B2523", on_click=next_question)

            elif step == "끝":
                if players:
                    ui.html(ceremony_html("개인"), sanitize=False).classes("contents")
                else:
                    ui.label("참가자가 없어요").classes("msg-big")
                with ui.element("div").classes("actions"):
                    if teams:
                        btn("👥 팀 시상식 보기", color="#8E44AD", on_click=show_team_ceremony)
                    btn("📋 전체 순위 보기", color="#D89E00", on_click=show_final_board)
                    btn("📊 결과 엑셀로 받기", color="#26890C", on_click=download_excel)
                    btn("새 게임 준비", color="#2B2523", on_click=reset_game)

            elif step == "팀끝":
                if team_rows():
                    ui.html(ceremony_html("팀"), sanitize=False).classes("contents")
                else:
                    ui.label("팀에 들어간 사람이 없어요").classes("msg-big")
                with ui.element("div").classes("actions"):
                    ui.button("← 개인 시상식", color="grey", on_click=back_to_ceremony)
                    btn("📋 전체 순위 보기", color="#D89E00", on_click=show_final_board)
                    btn("📊 결과 엑셀로 받기", color="#26890C", on_click=download_excel)
                    btn("새 게임 준비", color="#2B2523", on_click=reset_game)

            elif step == "최종표":
                with ui.element("div").classes("boards"):
                    draw_team_board("최종")
                    draw_board("최종", limit=BOARD_TOP_HOST)
                with ui.element("div").classes("actions"):
                    ui.button("← 개인 시상식", color="grey", on_click=back_to_ceremony)
                    if teams:
                        btn("👥 팀 시상식", color="#8E44AD", on_click=show_team_ceremony)
                    btn("📊 결과 엑셀로 받기", color="#26890C", on_click=download_excel)
                    btn("새 게임 준비", color="#2B2523", on_click=reset_game)

    def tick():
        check_time()
        key = screen_key()
        if key != shown["화면"]:
            shown["화면"] = key
            draw()
        # 소리: 장면(대기/문제/정답/순위/끝)이 바뀌면 음악과 효과음 바꾸기
        scene = (game["단계"], game["번호"])
        if scene != shown["장면"]:
            shown["장면"] = scene
            shown["카운트"] = None
            ui.run_javascript(f"quizSound.setScene('{game['단계']}')")
        # 마지막 5초는 띡, 띡, 띡
        if game["단계"] == "문제":
            left = time_left()
            if 0 < left <= 5 and left != shown["카운트"]:
                shown["카운트"] = left
                ui.run_javascript("quizSound.effect('tick')")
        if "시간" in live:
            answered = sum(1 for p in players.values() if p["답"] is not None)
            live["시간"].text = f"⏰ {time_left()}초"
            live["응답"].text = f"{answered} / {len(players)}명 답함"

    ui.timer(0.5, tick)   # 0.5초마다 칠판 확인


# =====================================================================
#  참가자 화면 (휴대폰)
# =====================================================================
@ui.page("/")
def player_page():
    saved = app.storage.user.get("닉네임")
    me = {"이름": saved if saved in players else None, "쫓겨남": False}
    area = ui.element("div").classes("screen")
    live = {}
    shown = {"화면": None}
    chosen = {"팀": None}     # 입장할 때 고른 팀

    def join(name):
        name = (name or "").strip()
        if not name:
            ui.notify("닉네임을 입력해주세요")
        elif teams and chosen["팀"] not in teams:
            ui.notify("팀을 골라주세요")
        elif name in kicked:
            ui.notify("진행자가 내보낸 닉네임이라 쓸 수 없어요")
        elif name in players:
            ui.notify("이미 있는 닉네임이에요. 다른 이름을 써주세요")
        elif game["단계"] != "대기":
            ui.notify("이미 게임이 시작됐어요. 다음 게임을 기다려주세요")
        else:
            players[name] = new_player(chosen["팀"] if teams else None)
            app.storage.user["닉네임"] = name
            me["이름"] = name
            me["쫓겨남"] = False

    def screen_key():
        if me["이름"] is None:
            return ("입장", me["쫓겨남"], tuple(teams))
        p = players[me["이름"]]
        return (game["단계"], game["번호"], p["답"] is not None, p["팀"], tuple(teams))

    def draw():
        area.clear()
        live.clear()
        with area:
            if me["이름"] is None:
                with ui.element("div").classes("join-box"):
                    if me["쫓겨남"]:
                        ui.label("진행자가 내보냈어요 🙅").classes("msg")
                    ui.label("닉네임을 입력하세요").classes("msg-big")
                    name_box = ui.input("닉네임", value=app.storage.user.get("닉네임", "")).props("outlined")
                    if teams:
                        ui.label("우리 팀을 골라주세요").classes("msg")
                        if chosen["팀"] not in teams:
                            chosen["팀"] = None
                        team_picker(lambda t: chosen.update(팀=t), selected=chosen["팀"])
                    btn("입장하기", color="#2B2523", on_click=lambda: join(name_box.value)
                              ).classes("text-lg")
                return

            name = me["이름"]
            p = players[name]
            step = game["단계"]

            if step == "대기":
                with ui.element("div").classes("join-box").style("align-items: center"):
                    ui.label(f"{name}님 입장 완료! 🙌").classes("msg-big")
                    if teams:
                        if p["팀"] in teams:
                            ui.label(f"우리 팀: {p['팀']}").classes("team-badge").style(
                                f"background: {team_color(p['팀'])}")
                            ui.label("시작 전까지는 팀을 바꿀 수 있어요").classes("msg")
                        else:
                            ui.label("⚠️ 팀을 골라주세요").classes("msg")

                        def change_team(t):
                            if game["단계"] == "대기" and t in teams:
                                p["팀"] = t
                        team_picker(change_team, selected=p["팀"])
                    ui.label("진행자가 시작하면 문제가 나와요").classes("msg")

            elif step == "문제" and p["답"] is None:
                draw_question_stage(
                    questions[game["번호"]], live,
                    on_pick=lambda n: submit_answer(me["이름"], n),
                    on_text=lambda text: submit_answer(me["이름"], text))

            elif step == "문제":
                ui.label("답을 냈어요! ✋").classes("msg-big")
                ui.label("다른 사람들을 기다리는 중...").classes("msg")

            elif step == "정답":
                q = questions[game["번호"]]
                if is_correct(q, p["답"]):
                    title = f"⭕ 정답! +{p['이번점수']}점"
                elif p["답"] is None:
                    title = "⏰ 시간 초과"
                else:
                    title = "❌ 아쉬워요"
                pills = [f"총점 {p['점수']}점"]
                if p["보너스"]:
                    pills.append(f"🔥 {p['연속']}연속 정답! +{p['보너스']}")
                draw_answer_stage(q, title, f"정답: {answer_text(q)}", pills, my_answer=p["답"])

            elif step in ("순위", "누적"):
                with ui.element("div").classes("boards"):
                    if step == "순위":
                        draw_board("라운드", my_name=name, limit=BOARD_TOP_PLAYER)
                        draw_team_board("라운드", my_team=p["팀"])
                    else:
                        draw_team_board("누적", my_team=p["팀"])
                        draw_board("누적", my_name=name, limit=BOARD_TOP_PLAYER)

            elif step == "팀끝":
                ui.html(ceremony_html("팀", my_team=p["팀"]), sanitize=False).classes("contents")

            elif step == "끝":
                ui.html(ceremony_html("개인", my_name=name), sanitize=False).classes("contents")
                team_order = [t for t, _, _ in team_rows()]
                if p["팀"] in team_order:
                    ui.label(f"우리 팀 {p['팀']}: {team_order.index(p['팀']) + 1}등").classes(
                        "team-badge").style(f"background: {team_color(p['팀'])}")

            elif step == "최종표":
                with ui.element("div").classes("boards"):
                    draw_team_board("최종", my_team=p["팀"])
                    draw_board("최종", my_name=name, limit=BOARD_TOP_PLAYER)

    def tick():
        if me["이름"] is not None and me["이름"] not in players:
            me["쫓겨남"] = me["이름"] in kicked
            me["이름"] = None
        check_time()
        key = screen_key()
        if key != shown["화면"]:
            shown["화면"] = key
            draw()
        if "시간" in live:
            live["시간"].text = f"⏰ {time_left()}초"

    ui.timer(0.5, tick)   # 0.5초마다 칠판 확인


# =====================================================================
#  문제 만들기 화면 (주소 끝에 /edit)
# =====================================================================
@ui.page("/edit")
def edit_page():
    if need_host_login():
        return
    photo = {"파일": ""}

    with ui.column().classes("w-full items-center gap-4 p-4"):
        ui.label("✏️ 문제 만들기").classes("text-4xl")
        ui.button("← 진행자 화면으로 돌아가기", color="grey",
                  on_click=lambda: ui.navigate.to("/host"))

        if game["단계"] != "대기":
            ui.label("게임 중에는 문제를 고칠 수 없어요. "
                     "게임이 끝나고 '새 게임 준비'를 누른 뒤 다시 와주세요.").classes("text-xl")
            return

        # --- 퀴즈 세트 고르기 / 만들기 / 지우기 ---
        with ui.card().classes("w-full max-w-2xl"):
            ui.label("📚 퀴즈 세트").classes("text-2xl")
            set_select = ui.select(list_sets(), value=game["세트"], label="지금 고른 세트",
                                   on_change=lambda e: change_set(e.value)).classes("w-64")
            with ui.row().classes("items-center"):
                new_set_box = ui.input("새 세트 이름 (예: 안전교육 퀴즈)").classes("w-64")
                ui.button("새 세트 만들기", on_click=lambda: create_set())
            ui.button("이 세트 지우기", color="red", on_click=lambda: confirm_delete_set()
                      ).props("outline")

        # --- 새 문제 입력칸 ---
        with ui.card().classes("w-full max-w-2xl"):
            ui.label("새 문제 추가").classes("text-2xl")
            type_select = ui.toggle(["4지선다", "OX", "단답형"], value="4지선다")
            q_box = ui.input("문제").classes("w-full")

            with ui.column().classes("w-full") as choice_area:
                choice_boxes = [ui.input(f"{n}번 보기").classes("w-full") for n in range(1, 5)]
                answer_select = ui.select({1: "1번", 2: "2번", 3: "3번", 4: "4번"},
                                          value=1, label="정답 번호").classes("w-40")
            choice_area.bind_visibility_from(type_select, "value", value="4지선다")

            with ui.column().classes("w-full") as ox_area:
                ui.label("정답")
                ox_select = ui.toggle({1: "⭕ O", 2: "❌ X"}, value=1)
            ox_area.bind_visibility_from(type_select, "value", value="OX")

            with ui.column().classes("w-full") as short_area:
                short_box = ui.input("정답 (여러 개면 쉼표로: 서울, 서울특별시)").classes("w-full")
                ui.label("띄어쓰기와 영어 대소문자는 달라도 정답으로 인정해요").classes("text-sm")
            short_area.bind_visibility_from(type_select, "value", value="단답형")

            ui.label("📷 사진 (없어도 돼요, 5MB 이하)")
            preview = ui.html("", sanitize=False)
            uploader = ui.upload(on_upload=lambda e: handle_upload(e), auto_upload=True,
                                 max_file_size=5_000_000,
                                 on_rejected=lambda: ui.notify("5MB 이하 사진만 올릴 수 있어요")
                                 ).props('accept="image/*"').classes("w-full")
            ui.button("추가하기", on_click=lambda: add())

        # --- 메모장 글로 한꺼번에 만들기 ---
        with ui.card().classes("w-full max-w-2xl"):
            ui.label("📝 메모장 글로 한꺼번에 만들기").classes("text-2xl")
            ui.label("메모장에 쓴 문제를 아래 칸에 붙여넣거나 .txt 파일을 올리면 문제로 바꿔줘요. "
                     "보기 4개 문제, OX, 단답형을 자동으로 알아봐요.").classes("text-sm")
            with ui.row().classes("items-center gap-2"):
                ui.button("예시 넣어보기", color="grey", on_click=lambda: fill_sample()).props("outline")
                ui.button("샘플 파일 받기", color="grey", on_click=lambda: download_sample()).props("outline")
            import_box = ui.textarea(
                placeholder="Q: 문제 글\n1) 보기1\n2) 보기2\n3) 보기3\n4) 보기4\n정답: 2"
            ).props("outlined autogrow").classes("w-full")
            ui.upload(label="또는 .txt 파일 올리기", on_upload=lambda e: handle_text_upload(e),
                      auto_upload=True, max_file_size=1_000_000,
                      on_rejected=lambda: ui.notify("1MB 이하 .txt 파일만 올릴 수 있어요")
                      ).props('accept=".txt,text/plain"').classes("w-full")
            mode_toggle = ui.toggle({"추가": "뒤에 추가", "교체": "기존 문제 모두 지우고 바꾸기"}, value="추가")
            ui.button("미리보기", color="#2B2523", on_click=lambda: preview_import()).props("text-color=white")
            import_result = ui.column().classes("w-full gap-1")

        list_area = ui.column().classes("w-full max-w-2xl")

    parsed = {"문제": []}

    def fill_sample():
        import_box.value = SAMPLE_TEXT
        preview_import()

    def download_sample():
        ui.download.content(SAMPLE_TEXT.encode("utf-8-sig"), "quiz_sample.txt")

    async def handle_text_upload(e):
        data = await e.file.read()
        import_box.value = decode_text(data)
        ui.notify(f"'{e.file.name}' 파일을 읽었어요")
        preview_import()

    def preview_import():
        """붙여넣은 글을 읽어서 어떻게 문제가 되는지 먼저 보여주기"""
        import_result.clear()
        good, errors = parse_quiz_text(import_box.value or "")
        parsed["문제"] = good
        with import_result:
            if not good and not errors:
                ui.label("읽을 글이 없어요. 위 칸에 문제를 붙여넣어주세요").classes("text-gray-600")
                return
            if good:
                ui.label(f"✅ 문제 {len(good)}개를 읽었어요").classes("text-lg")
                for i, q in enumerate(good, 1):
                    right = " / ".join(q["정답"]) if q["유형"] == "단답형" else q["보기"][q["정답"] - 1]
                    ui.label(f"{i}. [{q['유형']}] {q['문제']}  →  정답: {right}").classes("text-sm")
            for err in errors:
                ui.label(f"⚠️ {err}").classes("text-sm text-red-600")
            if errors:
                ui.label("⚠️ 위 문제들은 빼고 읽었어요. 글을 고치고 다시 미리보기를 눌러도 돼요.").classes("text-sm")
            if good:
                label = (f"✅ {len(good)}개 문제 넣기" if mode_toggle.value == "추가"
                         else f"⚠️ 기존 문제를 지우고 {len(good)}개로 바꾸기")
                ui.button(label, color="#26890C", on_click=lambda: apply_import()).props("text-color=white")

    def apply_import():
        if game["단계"] != "대기":
            ui.notify("게임 중에는 문제를 고칠 수 없어요")
            return
        new = parsed["문제"]
        if not new:
            return
        if mode_toggle.value == "교체":
            questions.clear()
        questions.extend(new)
        save_set()
        parsed["문제"] = []
        import_box.value = ""
        import_result.clear()
        ui.notify(f"문제 {len(new)}개를 넣었어요! 💾")
        draw_list()

    def draw_list():
        list_area.clear()
        with list_area:
            ui.label(f"'{game['세트']}' 세트 · 문제 {len(questions)}개").classes("text-2xl")
            for i, q in enumerate(questions):
                with ui.card().classes("w-full"):
                    ui.label(f"{i + 1}. [{q['유형']}] {q['문제']}").classes("text-xl")
                    if q["사진"]:
                        ui.html(photo_html(q, "q-photo-sm"), sanitize=False)
                    if q["유형"] == "단답형":
                        ui.label(f"정답: {' / '.join(q['정답'])}  ✅")
                    else:
                        for n, choice in enumerate(q["보기"], 1):
                            mark = "  ✅ 정답" if n == q["정답"] else ""
                            ui.label(f"{n}) {choice}{mark}")
                    ui.button("삭제", color="red", on_click=lambda i=i: delete(i))

    async def handle_upload(e):
        ext = Path(e.file.name).suffix.lower()
        if ext not in (".jpg", ".jpeg", ".png", ".gif", ".webp"):
            ui.notify("jpg, png, gif, webp 사진만 올릴 수 있어요")
            return
        filename = uuid.uuid4().hex + ext
        await e.file.save(PHOTO_DIR / filename)
        photo["파일"] = filename
        preview.content = (f'<img src="/images/questions/{filename}" '
                           f'style="max-height:160px; border-radius:10px">')
        ui.notify("사진을 올렸어요 📷")

    def add():
        if game["단계"] != "대기":
            ui.notify("게임 중에는 문제를 고칠 수 없어요")
            return
        text = (q_box.value or "").strip()
        if not text:
            ui.notify("문제를 입력해주세요")
            return
        kind = type_select.value
        if kind == "4지선다":
            choices = [(box.value or "").strip() for box in choice_boxes]
            if not all(choices):
                ui.notify("보기 4개를 모두 채워주세요")
                return
            q = {"유형": kind, "문제": text, "보기": choices, "정답": answer_select.value}
        elif kind == "OX":
            q = {"유형": kind, "문제": text, "보기": ["O", "X"], "정답": ox_select.value}
        else:
            answers = [a.strip() for a in (short_box.value or "").split(",") if a.strip()]
            if not answers:
                ui.notify("정답을 입력해주세요")
                return
            q = {"유형": kind, "문제": text, "보기": [], "정답": answers}
        q["사진"] = photo["파일"]
        questions.append(q)
        save_set()

        q_box.value = ""
        short_box.value = ""
        for box in choice_boxes:
            box.value = ""
        photo["파일"] = ""
        preview.content = ""
        uploader.reset()
        ui.notify("문제를 추가했어요! 💾")
        draw_list()

    def delete(i):
        if game["단계"] != "대기":
            ui.notify("게임 중에는 문제를 고칠 수 없어요")
            return
        del questions[i]
        save_set()
        ui.notify("삭제했어요")
        draw_list()

    def change_set(name):
        if choose_set(name):
            draw_list()

    def create_set():
        name = clean_name(new_set_box.value)
        if not name:
            ui.notify("세트 이름을 입력해주세요 (특수문자는 빼주세요)")
            return
        if name in list_sets():
            ui.notify("이미 있는 세트 이름이에요")
            return
        set_path(name).write_text("[]", encoding="utf-8")
        new_set_box.value = ""
        set_select.set_options(list_sets(), value=name)   # 새 세트로 바로 바뀌어요
        ui.notify(f"'{name}' 세트를 만들었어요")

    def confirm_delete_set():
        if len(list_sets()) <= 1:
            ui.notify("세트가 하나뿐이라 지울 수 없어요")
            return
        with ui.dialog() as dialog, ui.card():
            ui.label(f"'{game['세트']}' 세트를 정말 지울까요? 안의 문제도 모두 사라져요.")
            with ui.row():
                ui.button("지우기", color="red", on_click=lambda: (delete_set(), dialog.close()))
                ui.button("취소", on_click=dialog.close)
        dialog.open()

    def delete_set():
        set_path(game["세트"]).unlink(missing_ok=True)
        first = list_sets()[0]
        choose_set(first)
        set_select.set_options(list_sets(), value=first)
        draw_list()

    draw_list()


def port_in_use(port=8080):
    """8080번 문을 이미 다른 프로그램(예: 먼저 켜둔 퀴즈)이 쓰고 있는지"""
    import socket
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.settimeout(0.5)
        return sock.connect_ex(("127.0.0.1", port)) == 0


if port_in_use():
    print("\n" + "=" * 64)
    print("  [!] 퀴즈 프로그램이 이미 켜져 있어요 (8080번 문을 쓰는 중).")
    print("      IDLE 창이나 다른 검은 창으로 먼저 켜둔 퀴즈가 있는지 확인해서 닫아주세요.")
    print("      그 프로그램이 다른 폴더나 예전 버전이면, 만든 문제가 다르게 보일 수 있어요.")
    print("=" * 64 + "\n")
    raise SystemExit(1)

ui.run(host="0.0.0.0", reload=False, title="우리반 퀴즈",
       storage_secret="my-quiz-secret-1234")
